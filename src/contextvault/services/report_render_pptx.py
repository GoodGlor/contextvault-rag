"""Assemble a report as a PPTX deck (python-pptx).

The sibling of ``report_render.build_report_pdf``: same ``SectionResult`` input,
same ``report_format`` module, different container. Both artifacts are built at
generation time so downloading is a byte read and the reader picks the format.

A deck is not a paginated document, so the layout differs on purpose. ``build_report_pptx``
assembles the deck from the same ordered block list the document is built from, honoured at
**slide granularity** — which blocks exist, which section each binds to, and where slides
break. Within a slide the arrangement stays the deck's own: the statistics line renders as
the slide's caption, directly under the title, wherever ``stats`` sits in the block group
(the document prints it after the chart; the deck does not), and a chart consumes the rest
of its slide, so whatever follows it starts a continuation slide rather than being reflowed
onto the same one. Tables are capped harder than in the PDF — a slide that needs scrolling
is not a slide.

**``Block.icon`` is ignored here, deliberately (6d).** python-pptx cannot place inline SVG,
so a mark would have to be rasterised per occurrence — a rendering path, a font-colour
question and a cache, for a decoration. A report is not less itself without one. The
document draws marks and the deck does not, and that difference is asserted in
``tests/test_report_icons.py`` rather than left to this paragraph, because the field rides
on the same block list both renderers read.
"""

import struct
import uuid
import zipfile
from collections.abc import Callable, Mapping
from dataclasses import dataclass
from io import BytesIO

from pptx import Presentation as new_presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.presentation import Presentation
from pptx.slide import Slide
from pptx.util import Inches, Length, Pt

from contextvault.services.report_blocks import Block, default_blocks, document_title
from contextvault.services.report_format import (
    format_value,
    infer_column_kind,
    label,
    measure_summary,
    rows_note,
    short_error,
)
from contextvault.services.report_style import ReportStyle, resolve_style
from contextvault.services.report_types import ReportAssetRef, SectionResult

# A 16:9 slide fits far less than an A4 page; beyond this the text is unreadable.
_MAX_TABLE_ROWS = 12
_MAX_TABLE_COLS = 6
_MAX_STAT_COLUMNS = 3  # one line of stats; a slide has no room for more

_SLIDE_W = Inches(13.333)
_SLIDE_H = Inches(7.5)
_MARGIN = Inches(0.6)
_CONTENT_W = Length(Inches(13.333) - 2 * Inches(0.6))
_BODY_TOP = Inches(1.5)

# python-pptx writes each zip entry with a wall-clock timestamp, so two builds of
# one deck differ in bytes while being identical in every part. A fixed stamp makes
# a rebuild genuinely reproduce its artifact rather than merely resemble it — the
# same reason step 5a pinned matplotlib's SVG date and hash salt. 1980-01-01 is the
# zero of the DOS timestamp a zip entry stores; only its constancy matters.
_ZIP_EPOCH = (1980, 1, 1, 0, 0, 0)

# The three zip record signatures this module reaches for, and where each keeps
# the DOS time+date pair: 10 bytes into a local file header, 12 into a central
# directory entry. Both offsets are fixed by the format.
_LOCAL_HEADER = b"PK\x03\x04"
_CENTRAL_ENTRY = b"PK\x01\x02"
_END_OF_CENTRAL_DIR = b"PK\x05\x06"
_LOCAL_STAMP_AT = 10
_CENTRAL_STAMP_AT = 12


def _rgb(colour: str) -> RGBColor:
    """A ``#rrggbb`` token as python-pptx's colour type.

    python-pptx leaves both ``RGBColor`` and its ``from_string`` untyped, so the
    ignore is unavoidable; it sits here, once, rather than at every call site.
    """
    red, green, blue = (int(colour.lstrip("#")[index : index + 2], 16) for index in (0, 2, 4))
    return RGBColor(red, green, blue)  # type: ignore[no-untyped-call]


def _blank(presentation: Presentation, style: ReportStyle) -> Slide:
    """A slide with no placeholders — every element is positioned explicitly.

    A deck has no stylesheet, so nothing propagates and nothing is inherited: the
    ground is painted here, on every slide, because a slide added anywhere else
    comes back white.
    """
    slide: Slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = _rgb(style.paper)
    return slide


def _textbox(
    slide: Slide,
    text: str,
    *,
    top: Length,
    size: int,
    style: ReportStyle,
    bold: bool = False,
) -> Length:
    """Draw a wrapped paragraph and return the vertical position just below it."""
    width = _CONTENT_W
    box = slide.shapes.add_textbox(_MARGIN, top, width, Inches(0.8))
    frame = box.text_frame
    frame.word_wrap = True
    frame.text = text
    for paragraph in frame.paragraphs:
        for run in paragraph.runs:
            run.font.size = Pt(size)
            run.font.bold = bold
            # Painting the ground and leaving the type black is worse than doing
            # neither: the deck would be unreadable rather than merely plain.
            run.font.color.rgb = _rgb(style.ink)
    # python-pptx cannot measure wrapped text, so estimate from character count.
    lines = max(1, len(text) // max(1, int(90 * 18 / max(size, 1))) + 1)
    return Length(top + Pt(size * 1.35 * lines) + Inches(0.15))


def _dos_stamp(when: tuple[int, int, int, int, int, int]) -> bytes:
    """A date as the packed DOS time+date pair a zip header carries.

    Seconds have two-second resolution in this encoding, which is why the epoch
    above sits on zero rather than relying on the field to round.
    """
    year, month, day, hour, minute, second = when
    return struct.pack(
        "<HH",
        (hour << 11) | (minute << 5) | (second // 2),
        ((year - 1980) << 9) | (month << 5) | day,
    )


_ZIP_STAMP = _dos_stamp(_ZIP_EPOCH)


def _with_fixed_timestamps(package: bytes) -> bytes:
    """The same package, every zip entry stamped ``_ZIP_EPOCH``.

    Rewrites headers only, and means it literally: the two date/time fields of
    each local header and of its central-directory entry are overwritten in
    place, and every other byte is the one python-pptx wrote. Nothing is
    inflated and nothing re-deflated, so no part's content can move — and the
    goldens do not quietly depend on which zlib build compressed them.

    ``infolist`` gives the local headers their offsets; it does not say where
    each central-directory entry sits, so those are walked from the
    end-of-central-directory record, which names the directory's start, size and
    entry count. The closing assertion is the check that the walk landed exactly
    on the directory's end rather than wandering through it.
    """
    buffer = bytearray(package)
    for info in zipfile.ZipFile(BytesIO(package)).infolist():
        at = info.header_offset + _LOCAL_STAMP_AT
        assert buffer[info.header_offset : info.header_offset + 4] == _LOCAL_HEADER
        buffer[at : at + 4] = _ZIP_STAMP

    end = package.rfind(_END_OF_CENTRAL_DIR)
    assert end != -1, "not a zip package"
    count, size, start = struct.unpack("<HII", package[end + 10 : end + 20])
    at = start
    for _ in range(count):
        assert buffer[at : at + 4] == _CENTRAL_ENTRY
        buffer[at + _CENTRAL_STAMP_AT : at + _CENTRAL_STAMP_AT + 4] = _ZIP_STAMP
        name, extra, comment = struct.unpack("<HHH", package[at + 28 : at + 34])
        at += 46 + name + extra + comment
    assert at == start + size
    return bytes(buffer)


def _title_slide(presentation: Presentation, title: str, style: ReportStyle) -> None:
    slide = _blank(presentation, style)
    _textbox(slide, title, top=Inches(2.6), size=40, style=style, bold=True)


def _stats_line(section: SectionResult, language: str) -> str:
    """The same measures the PDF prints, as one line — a slide has no room for a list.

    "The same" is load-bearing and was not true. This used to be a parallel
    implementation of ``report_html._numeric_stats`` and it drifted twice, each
    time by learning a rule the other path already knew: #176's ``keys_repeat``
    and #178's ``overlap_warning``. Both now live once, in
    ``report_format.measure_summary``, which this and the document share — so a
    rule can no longer reach one artifact and miss the other.

    What stays here is the medium: ``_MAX_STAT_COLUMNS`` is 3 because a slide has
    no room for more, and the pairs are joined into a single line rather than
    laid out as a list.
    """
    result = section.result
    assert result is not None
    stats = measure_summary(
        result.columns,
        result.rows,
        language=language,
        column_labels=section.column_labels or {},
        max_columns=_MAX_STAT_COLUMNS,
        overlapping=section.overlap_warning is not None,
    )
    return "   ·   ".join(f"{key}: {value}" for key, value in stats)


def _add_table(
    slide: Slide, section: SectionResult, language: str, top: Length, style: ReportStyle
) -> None:
    result = section.result
    assert result is not None
    labels = section.column_labels or {}
    columns = result.columns[:_MAX_TABLE_COLS]
    kinds = [
        infer_column_kind(name, [row[i] for row in result.rows])
        for i, name in enumerate(result.columns)
    ][:_MAX_TABLE_COLS]
    rows = result.rows[:_MAX_TABLE_ROWS]
    if not rows:
        return
    shape = slide.shapes.add_table(
        len(rows) + 1, len(columns), _MARGIN, top, _CONTENT_W, Inches(0.4)
    )
    table = shape.table
    for col_index, name in enumerate(columns):
        cell = table.cell(0, col_index)
        cell.text = str(labels.get(name, name))
        for paragraph in cell.text_frame.paragraphs:
            for run in paragraph.runs:
                run.font.size = Pt(11)
                run.font.bold = True
                run.font.color.rgb = _rgb(style.ink)
    for row_index, row in enumerate(rows, start=1):
        for col_index in range(len(columns)):
            cell = table.cell(row_index, col_index)
            cell.text = format_value(row[col_index], kinds[col_index], language)
            for paragraph in cell.text_frame.paragraphs:
                for run in paragraph.runs:
                    run.font.size = Pt(10)
                    run.font.color.rgb = _rgb(style.ink)


@dataclass(frozen=True)
class _SlideGroup:
    """One slide's worth of blocks, plus how its title is drawn.

    ``heading`` distinguishes the slide a ``heading`` block opens — the section
    title at 26 pt — from a continuation slide at 20 pt, which is the same
    distinction ``_data_slides`` drew inline before the deck became block-driven.

    ``title`` is the words that slide is titled with: the heading block's override
    where there is one, else the section's own. It is carried here rather than
    read from the heading at draw time because a *continuation* group has no
    heading block — it would otherwise print the generated title beside a first
    slide wearing the reader's, naming one section two ways in one deck.
    """

    section: SectionResult | None
    heading: bool
    blocks: tuple[Block, ...]
    title: str | None = None


def _section_of(block: Block, sections: list[SectionResult]) -> SectionResult | None:
    """The section a block binds to, or ``None`` when the index does not resolve.

    ``None`` is the read-tolerance signal, not an error: a stored list whose
    indexes have gone stale must degrade to silence, never to an exception.
    """
    if block.section is None or not 0 <= block.section < len(sections):
        return None
    return sections[block.section]


def _agenda_entries(blocks: list[Block], sections: list[SectionResult]) -> tuple[str, ...]:
    """The section titles an agenda slide lists.

    The reader's 5e heading override wins over the model's title, for the same
    reason it wins on the slide itself: the agenda must name a section the way
    the section names itself. A hidden heading opens no slide, so listing it
    would point at something the deck does not contain.
    """
    titles: list[str] = []
    for block in blocks:
        if block.kind != "heading" or block.hidden:
            continue
        bound = _section_of(block, sections)
        if bound is None:
            continue
        titles.append(block.text if block.text is not None else bound.title)
    return tuple(titles)


def _group_slides(blocks: list[Block], sections: list[SectionResult]) -> list[_SlideGroup]:
    """Split the block list into slides.

    ``heading`` opens a titled slide, ``pagebreak`` opens a continuation of the
    section already open, and ``cover`` is a slide of its own. A body block
    arriving with no slide open starts an untitled one rather than being dropped
    — silence is for blocks that cannot be rendered, not for blocks that arrive
    early. A ``pagebreak`` before any heading has nothing to continue, so it is a
    no-op: an empty slide is a visible defect and a dropped break is not.
    """
    groups: list[_SlideGroup] = []
    pending: list[Block] = []
    section: SectionResult | None = None
    section_title: str | None = None
    heading = False

    def close() -> None:
        nonlocal pending
        if pending:
            groups.append(_SlideGroup(section, heading, tuple(pending), section_title))
            pending = []

    for block in blocks:
        # Skipped before grouping — the single point that decides what reaches a
        # slide — so a hidden block neither draws nor reserves space, and a slide
        # whose every block is hidden is not produced at all.
        if block.hidden:
            continue
        if block.kind == "cover":
            close()
            groups.append(_SlideGroup(None, False, (block,)))
            # A mid-list cover ends whatever section was open. Leaving it bound
            # would make the next body block re-render as though a fresh heading
            # had just opened it — the same title, a second time, at 26 pt.
            section, section_title, heading = None, None, False
            continue
        if block.kind == "toc":
            # Its own slide, like the cover: an agenda shares a slide with
            # nothing, and it binds to no section.
            close()
            groups.append(_SlideGroup(None, False, (block,)))
            continue
        if block.kind == "heading":
            bound = _section_of(block, sections)
            if bound is None:
                continue
            close()
            section, heading = bound, True
            section_title = block.text if block.text is not None else bound.title
            pending = [block]
            continue
        if block.kind == "pagebreak":
            if section is None and not pending:
                continue
            close()
            heading = False
            continue
        pending.append(block)
    close()
    return groups


_Canvas = Callable[[], tuple[Slide, Length]]
"""What a block asks for when it has something to draw: a slide, and the cursor on it.

The slide behind a canvas is created — and titled — by the *first* call, so a
block that never asks costs nothing. That is the whole blank-slide guard: rather
than predicting whether a group would draw ("would this draw" being a second
opinion free to drift from ``_draw_block``'s own branches, and blind to the
continuation slide a chart opens), slide allocation simply follows the shapes.
"""


def _draw_callout(slide: Slide, text: str, *, top: Length, style: ReportStyle) -> Length:
    """A tinted band — the deck's answer to the document's ``.callout`` CSS rule.

    A deck has no stylesheet, so the emphasis is drawn: the template's ``band``
    token, the same fill that sits behind a table heading.
    """
    height = Inches(0.9)
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, _MARGIN, top, _CONTENT_W, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = _rgb(style.band)
    shape.line.fill.background()
    frame = shape.text_frame
    frame.word_wrap = True
    frame.text = text
    for paragraph in frame.paragraphs:
        for run in paragraph.runs:
            run.font.size = Pt(14)
            run.font.color.rgb = _rgb(style.ink)
    return Length(top + height + Inches(0.15))


def _draw_table_block(
    canvas: _Canvas, section: SectionResult, *, language: str, style: ReportStyle
) -> Length | None:
    """The section's rows, or the note that says there are none.

    The empty case is the table block's job, exactly as it is in the document —
    ``default_blocks`` derives no chart and no statistics line for an empty
    result, so this note is the only thing the section renders. A section with no
    ``result`` at all is the one case with nothing to say, so it asks for no
    canvas and reports ``None``.

    The overlap caveat (#178) is drawn here, above the rows, for the reason the
    document draws it in ``_table_body``: it is a statement about *these rows*,
    and a reader meets it with the thing it is about. With no rows there is
    nothing being double-counted and nothing to caveat — the same rule the
    document follows, and the reason the deck goldens do not move.
    """
    result = section.result
    if result is None:
        return None
    slide, top = canvas()
    if not result.rows:
        top = _textbox(slide, label("no_rows", language), top=top, size=14, style=style)
        for note in section.notes:
            top = _textbox(slide, note, top=top, size=11, style=style)
        return top
    if section.overlap_warning is not None:
        top = _textbox(slide, label("overlapping_periods", language), top=top, size=11, style=style)
    # #204: which databases contributed nothing, above the rows for the same
    # reason the caveat is — it is a statement about the table underneath it.
    for note in section.notes:
        top = _textbox(slide, note, top=top, size=11, style=style)
    _add_table(slide, section, language, top, style)
    omitted = len(result.rows) - _MAX_TABLE_ROWS
    if omitted > 0:
        _textbox(
            slide, rows_note("more_rows", omitted, language), top=Inches(6.7), size=11, style=style
        )
    # The row height ``_add_table`` itself passes, header row included — a table
    # was always the last block on its slide before this task, so nothing read
    # this return value; a reordered list can now put a callout or a second
    # table after it, and both would otherwise land on top of the rows.
    return Length(top + Inches(0.4) * (min(len(result.rows), _MAX_TABLE_ROWS) + 1) + Inches(0.15))


def _draw_image_block(
    canvas: _Canvas, block: Block, assets: Mapping[uuid.UUID, ReportAssetRef]
) -> Length | None:
    """An uploaded image, scaled to the block's width and kept in proportion.

    Only the width is given to ``add_picture``: python-pptx then derives the
    height from the file's own pixels, so a logo cannot be stretched by a block
    whose percentage does not match its aspect ratio.

    A design can outlive its picture — assets are deleted explicitly — so an
    unresolvable reference draws nothing and asks for no slide, rather than
    leaving an empty frame on a slide of its own.
    """
    asset = assets.get(block.asset_id) if block.asset_id is not None else None
    if asset is None:
        return None
    percent = block.width if block.width is not None else 70
    width = Length(int(_CONTENT_W * percent / 100))
    slide, top = canvas()
    picture = slide.shapes.add_picture(BytesIO(asset.data), _MARGIN, top, width=width)
    return Length(top + picture.height + Inches(0.2))


def _draw_block(
    canvas: _Canvas,
    block: Block,
    section: SectionResult | None,
    *,
    language: str,
    style: ReportStyle,
    assets: Mapping[uuid.UUID, ReportAssetRef],
    agenda: tuple[str, ...],
) -> Length | None:
    """One block, on the slide it asks ``canvas`` for; the cursor just below it.

    ``None`` means the block drew nothing and therefore asked for no slide: an
    unknown kind, a block whose section did not resolve, a citations line for a
    section with no sources, a table over a section with no result. That is how
    the deck stays renderable when a stored list outlives the code or the
    sections it names — and how the silence stays silent, rather than
    materialising as a slide with nothing on it.
    """
    if block.kind == "callout":
        slide, top = canvas()
        return _draw_callout(slide, block.text or "", top=top, style=style)
    if block.kind == "toc":
        # An agenda, not a contents page: a deck has no pages, so there is no
        # number to print beside an entry, and inventing slide indexes would be
        # a second numbering nobody asked for.
        if not agenda:
            return None
        slide, top = canvas()
        cursor = _textbox(slide, label("contents", language), top=top, size=20, style=style)
        return _textbox(slide, "\n".join(agenda), top=cursor, size=14, style=style)
    if block.kind == "image":
        return _draw_image_block(canvas, block, assets)
    if section is None:
        return None
    if block.kind == "prose":
        slide, top = canvas()
        words = block.text if block.text is not None else section.prose or ""
        return _textbox(slide, words, top=top, size=14, style=style)
    if block.kind == "sources":
        # An override draws even where the section resolved no citations — the
        # reader wrote a line, and dropping it would read as the edit being lost.
        if block.text is None and not section.sources:
            return None
        listed = (
            block.text
            if block.text is not None
            else " · ".join(f"[{ref.number}] {ref.title}" for ref in section.sources)
        )
        line = "Джерела / Sources: " + listed
        slide, top = canvas()
        return _textbox(slide, line, top=top, size=10, style=style)
    if block.kind == "error":
        slide, top = canvas()
        cursor = _textbox(
            slide, f"⚠ {label('section_failed', language)}", top=top, size=14, style=style
        )
        return _textbox(slide, short_error(section.error), top=cursor, size=10, style=style)
    if block.kind == "viz" and block.variant == "table":
        return _draw_table_block(canvas, section, language=language, style=style)
    return None


def _draw_group(
    presentation: Presentation,
    group: _SlideGroup,
    sections: list[SectionResult],
    *,
    title: str,
    language: str,
    style: ReportStyle,
    assets: Mapping[uuid.UUID, ReportAssetRef],
    agenda: tuple[str, ...],
) -> None:
    """Draw one group as slides, or as the title slide when it is the cover.

    Two rules here are the deck's own rather than the list's, and both predate
    it. The statistics line is the slide's **caption**: it sits under the title
    wherever the list puts it, because a slide has one caption slot — and for the
    same reason a *second* ``stats`` block in the group is dropped, where the
    document would print a second statistics line. And a chart consumes the rest
    of its slide, so whatever follows it starts a continuation.

    Every slide here is created lazily, through ``canvas``, by the first block
    with something to put on it — including a ``heading``, whose shape is the
    slide title itself. A group whose blocks all decline — stale indexes, a
    chart with no image, a citations line with no sources — therefore allocates
    no slide at all, on the first slide and on any continuation alike.
    """
    if group.blocks and group.blocks[0].kind == "cover":
        _title_slide(presentation, title, style)
        return

    slide: Slide | None = None
    cursor: Length = _MARGIN
    opened = 0

    def canvas() -> tuple[Slide, Length]:
        """The group's current slide and cursor, opening the slide on first use."""
        nonlocal slide, cursor, opened
        if slide is None:
            slide = _blank(presentation, style)
            cursor = _MARGIN
            if group.section is not None:
                # 26 pt is the section title a ``heading`` opens; a continuation
                # repeats it at 20 pt, which is the distinction ``_data_slides``
                # drew inline before the deck became block-driven.
                cursor = _textbox(
                    slide,
                    group.title if group.title is not None else group.section.title,
                    top=cursor,
                    size=26 if group.heading and opened == 0 else 20,
                    style=style,
                    bold=True,
                )
            opened += 1
        return slide, cursor

    stats = next((b for b in group.blocks if b.kind == "stats"), None)
    if stats is not None:
        bound = _section_of(stats, sections)
        if bound is not None and bound.result is not None and bound.result.rows:
            target, top = canvas()
            cursor = _textbox(target, _stats_line(bound, language), top=top, size=12, style=style)

    for block in group.blocks:
        if block.kind == "stats":
            continue  # already drawn, as the caption, above
        if block.kind == "heading":
            # The one block whose shape *is* the slide title, which ``canvas``
            # draws. So a heading opens its slide even when nothing follows it:
            # a titled slide is not a blank one, and dropping it would lose the
            # slide a ``pagebreak`` immediately after a heading asks for.
            canvas()
            continue
        section = _section_of(block, sections)
        if block.kind == "viz" and block.variant == "chart":
            # The chart is placed here rather than in ``_draw_block`` because it
            # alone ends its slide; no picture behind the block means no slide.
            if section is None or section.chart_png is None:
                continue
            target, top = canvas()
            target.shapes.add_picture(BytesIO(section.chart_png), _MARGIN, top, width=_CONTENT_W)
            if len(section.databases) > 1:
                # #205: the document captions the figure inside its ``<figure>``;
                # a deck has no stylesheet, so the same statement becomes a line
                # under the picture. Placed at a fixed foot rather than measured
                # from the image, the way the omitted-rows note already is — the
                # picture's drawn height is not known here.
                _textbox(
                    target,
                    label("drawn_from", language).format(names=", ".join(section.databases)),
                    top=Inches(6.7),
                    size=11,
                    style=style,
                )
            slide = None  # a chart consumes the rest of its slide
            continue
        drawn = _draw_block(
            canvas,
            block,
            section,
            language=language,
            style=style,
            assets=assets,
            agenda=agenda,
        )
        if drawn is not None:
            cursor = drawn


def build_report_pptx(
    *,
    title: str,
    sections: list[SectionResult],
    language: str = "en",
    style: ReportStyle | None = None,
    blocks: list[Block] | None = None,
    assets: Mapping[uuid.UUID, ReportAssetRef] | None = None,
) -> bytes:
    """The blended report as a deck, assembled from an ordered block list.

    ``blocks`` is the same list the document is built from; ``None`` means the
    report is unedited and the default is derived. The list is honoured at slide
    granularity — which blocks exist, which section each binds to, and where
    slides break — while the arrangement within a slide stays the deck's, which
    is what keeps an unedited deck identical to the one this renderer produced
    before it read blocks at all.
    """
    style = style or resolve_style(None)
    presentation = new_presentation()
    presentation.slide_width = _SLIDE_W
    presentation.slide_height = _SLIDE_H

    block_list = blocks if blocks is not None else default_blocks(sections)
    # The same helper the document's <title> uses, so the two artifacts cannot
    # end up naming one report differently.
    name = document_title(title, block_list)
    agenda = _agenda_entries(block_list, sections)

    for group in _group_slides(block_list, sections):
        _draw_group(
            presentation,
            group,
            sections,
            title=name,
            language=language,
            style=style,
            assets=assets or {},
            agenda=agenda,
        )

    buffer = BytesIO()
    presentation.save(buffer)
    return _with_fixed_timestamps(buffer.getvalue())
