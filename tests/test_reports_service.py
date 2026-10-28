"""Orchestration: plan → per-section execute → assemble, self-repair, partial success."""

import json
import uuid
from collections.abc import AsyncIterator
from contextlib import asynccontextmanager
from datetime import UTC, date, datetime, time, timedelta
from decimal import Decimal
from typing import Any

import pytest
import sqlalchemy as sa
from sqlalchemy.engine.url import make_url
from sqlalchemy.ext.asyncio import AsyncSession

import contextvault.services.report_planner as report_planner
import contextvault.services.reports as reports_service
from contextvault.core.config import get_settings
from contextvault.core.crypto import encrypt
from contextvault.models import (
    DatabaseConnection,
    DatabaseType,
    GeneratedReport,
    LLMProviderName,
    ProviderSetting,
    ReportSchedule,
    ReportStatus,
    Repository,
    Source,
    SourceKind,
    SourceStatus,
)
from contextvault.services.report_blocks import BLOCKS_VERSION
from contextvault.services.report_design import design_payload
from contextvault.services.report_execution import QueryResult
from contextvault.services.report_format import label
from contextvault.services.report_llm import ChartSpec, ReportQuery
from contextvault.services.report_persist import load_sections
from contextvault.services.report_planner import (
    PROVENANCE_COLUMN,
    DraftSection,
    PlannedSection,
    ReportPlan,
    ReportPlanDraft,
    ReportPlanParseError,
)
from contextvault.services.report_style import TEMPLATES, resolve_style
from contextvault.services.report_types import SectionResult
from contextvault.services.reports import (
    _MAX_PAYLOAD_BYTES,
    ReportNotRenderable,
    _store_sections,
    preview_report_html,
    rerender_report,
)


def _fixed_factory(session: AsyncSession):  # type: ignore[no-untyped-def]
    """A session factory that always yields ``session`` without closing it."""

    @asynccontextmanager
    async def factory() -> AsyncIterator[AsyncSession]:
        yield session

    return factory


@pytest.fixture
def report_row() -> GeneratedReport:
    # Not persisted: _store_sections only assigns an attribute.
    return GeneratedReport(prompt="Звіт", repository_id=uuid.uuid4())


@pytest.fixture
def one_data_section() -> SectionResult:
    return SectionResult(
        title="Виторг",
        source="database",
        sub_question="Скільки?",
        kind="data",
        result=QueryResult(columns=["fy", "revenue"], rows=[(2024, Decimal("391.0"))]),
        chart_spec={"chart_type": "bar", "x_column": "fy", "y_column": "revenue"},
    )


@pytest.fixture
def huge_data_section() -> SectionResult:
    rows = [(year, Decimal("1.5")) for year in range(400_000)]
    return SectionResult(
        title="Забагато",
        source="database",
        sub_question="?",
        kind="data",
        result=QueryResult(columns=["n", "value"], rows=rows),
    )


def test_payload_is_stored_with_title_and_language(
    report_row: GeneratedReport, one_data_section: SectionResult
) -> None:
    _store_sections(report_row, title="Звіт", language="uk", sections=[one_data_section])
    assert report_row.sections_data is not None
    assert report_row.sections_data["title"] == "Звіт"
    assert report_row.sections_data["language"] == "uk"
    assert len(report_row.sections_data["sections"]) == 1


def test_an_oversized_payload_is_refused_rather_than_truncated(
    report_row: GeneratedReport, huge_data_section: SectionResult
) -> None:
    # Truncating would re-render a different chart; the report reports itself
    # unrenderable instead.
    _store_sections(report_row, title="Звіт", language="uk", sections=[huge_data_section])
    assert report_row.sections_data is None


def test_the_budget_is_measured_on_the_serialized_payload(
    report_row: GeneratedReport, one_data_section: SectionResult
) -> None:
    _store_sections(report_row, title="Звіт", language="uk", sections=[one_data_section])
    size = len(json.dumps(report_row.sections_data, ensure_ascii=False).encode())
    assert size <= _MAX_PAYLOAD_BYTES


def _single_db_plan(sub_question: str, conn: DatabaseConnection) -> dict[str, Any]:
    """The stored plan a one-database-section report generates from.

    Since #203 a database section names the connection it reads, so a stored plan
    is only executable against the repository it was planned for.
    """
    return {
        "title": sub_question,
        "language": "",
        "sections": [
            {
                "title": sub_question,
                "source": "database",
                "databases": [str(conn.id)],
                "sub_question": sub_question,
            }
        ],
    }


async def _connected_repo(db_session: AsyncSession) -> tuple[Repository, DatabaseConnection]:
    url = make_url(get_settings().database_url)
    repo = Repository(name="Vault", llm_provider="gemini", llm_model="gem-1")
    db_session.add(repo)
    # ``generate_report`` gates on ``repo_is_answerable``, which requires a verified
    # provider setting (not just the monkeypatched key lookup below).
    db_session.add(
        ProviderSetting(
            provider=LLMProviderName.GEMINI,
            api_key_encrypted=encrypt("unused-real-key"),
            verified_at=datetime(2026, 1, 1, tzinfo=UTC),
        )
    )
    await db_session.flush()
    conn = DatabaseConnection(
        repository_id=repo.id,
        name="Reporting",
        db_type=DatabaseType.POSTGRES,
        host=url.host or "localhost",
        port=url.port or 5432,
        database=url.database or "",
        username=url.username or "",
        password_encrypted=encrypt(url.password or ""),
        exposed_schema=[
            {
                "table": "repositories",
                "description": "",
                "columns": [{"name": "name", "description": ""}],
            }
        ],
    )
    db_session.add(conn)
    await db_session.flush()
    return repo, conn


async def _docs_only_repo(db_session: AsyncSession) -> Repository:
    """A repository whose only capability is an ingested document."""
    repo = Repository(name="Docs vault", llm_provider="gemini", llm_model="gem-1")
    db_session.add(repo)
    db_session.add(
        ProviderSetting(
            provider=LLMProviderName.GEMINI,
            api_key_encrypted=encrypt("unused-real-key"),
            verified_at=datetime(2026, 1, 1, tzinfo=UTC),
        )
    )
    await db_session.flush()
    db_session.add(
        Source(
            repository_id=repo.id,
            title="policy.pdf",
            kind=SourceKind.DOCUMENT,
            status=SourceStatus.DONE,
        )
    )
    await db_session.flush()
    return repo


def _mixed_plan(conn: DatabaseConnection) -> dict[str, Any]:
    """The stored plan for a report whose sections span both sources."""
    return {
        "title": "",
        "language": "",
        "sections": [
            {
                "title": "Числа",
                "source": "database",
                "databases": [str(conn.id)],
                "sub_question": "names",
            },
            {"title": "Політика", "source": "documents", "sub_question": "правила"},
        ],
    }


async def _matching_rows(**kwargs: Any) -> QueryResult:
    """A query that matched something, whatever the database happens to hold.

    These tests execute their SQL for real, and ``run_validated_query`` opens its
    **own** connection — so the repository ``_connected_repo`` inserted through
    ``db_session`` is not visible to it, and the query returns whatever earlier
    runs left committed. On a clean database that is nothing at all.

    That was invisible while a report with no rows was still delivered as a
    success. Since #185 it refuses the report, so a test that meant to exercise
    rendering or the plan path would fail on an empty database and pass on a
    dirty one — which is worse than either. Filtered queries (``WHERE name =
    'Apple'``) still match nothing, because that is what they are testing.
    """
    sql = str(kwargs.get("sql", ""))
    if "'Apple'" in sql or "Tesla" in sql:
        return QueryResult(columns=["name"], rows=[])
    return QueryResult(columns=["name"], rows=[("Бухгалтерія",), ("Склад",)])


def _query(sql: str) -> ReportQuery:
    return ReportQuery(sql=sql, chart=ChartSpec(chart_type="none", title="t"))


async def test_self_repair_recovers_and_records_final_sql(
    db_session: AsyncSession, monkeypatch: pytest.MonkeyPatch
) -> None:
    repo, conn = await _connected_repo(db_session)
    report = GeneratedReport(
        repository_id=repo.id,
        connection_id=conn.id,
        requested_by=None,
        prompt="names",
        plan=_single_db_plan("names", conn),
    )
    db_session.add(report)
    await db_session.flush()

    attempts: list[list[str]] = []
    # The recovering query must return a row on ANY database, so it counts rather
    # than selects. ``SELECT name FROM repositories`` passed only because a
    # developer's database happens to hold repositories; against a freshly
    # migrated one it returns nothing, and #133's "a query matching nothing is not
    # a success" rule then spends a third self-repair attempt that this test never
    # supplies a reply for — so the section failed, the report failed, and the
    # assertion below blamed self-repair. ``count(*)`` returns exactly one row
    # whether the table holds 165 rows or none.
    replies = iter(
        [_query("SELECT secret FROM vault_x"), _query("SELECT count(*) AS name FROM repositories")]
    )

    async def fake_generate(provider: str, api_key: str, model: str, **kwargs: Any) -> ReportQuery:
        attempts.append(list(kwargs["feedback"]))
        return next(replies)

    monkeypatch.setattr(reports_service, "generate_report_query", fake_generate)
    monkeypatch.setattr(
        reports_service.provider_service,  # type: ignore[attr-defined]
        "get_provider_key",
        _fake_key,
    )

    await reports_service.generate_report(db_session, report)
    assert report.status is ReportStatus.DONE
    assert report.generated_sql is None
    assert report.sections is not None
    assert report.sections[0]["generated_sql"] is not None
    assert "repositories" in report.sections[0]["generated_sql"]
    assert report.pdf_data is not None and report.pdf_data[:5] == b"%PDF-"
    assert len(attempts) == 2 and attempts[1]  # second attempt carried feedback


async def _fake_key(session: AsyncSession, provider: LLMProviderName) -> str:
    return "k"


async def test_three_bad_attempts_fail_the_report(
    db_session: AsyncSession, monkeypatch: pytest.MonkeyPatch
) -> None:
    repo, conn = await _connected_repo(db_session)
    report = GeneratedReport(
        repository_id=repo.id,
        connection_id=conn.id,
        requested_by=None,
        prompt="p",
        plan=_single_db_plan("p", conn),
    )
    db_session.add(report)
    await db_session.flush()

    async def always_bad(provider: str, api_key: str, model: str, **kwargs: Any) -> ReportQuery:
        return _query("DROP TABLE repositories")

    monkeypatch.setattr(reports_service, "generate_report_query", always_bad)
    monkeypatch.setattr(
        reports_service.provider_service,  # type: ignore[attr-defined]
        "get_provider_key",
        _fake_key,
    )

    await reports_service.generate_report(db_session, report)
    assert report.status is ReportStatus.FAILED
    assert report.error is not None and "rejected" in report.error


async def test_blended_report_mixes_db_and_documents_sections(
    db_session: AsyncSession, monkeypatch: pytest.MonkeyPatch
) -> None:
    repo, conn = await _connected_repo(db_session)
    report = GeneratedReport(
        repository_id=repo.id,
        connection_id=conn.id,
        requested_by=None,
        prompt="огляд",
        plan=_mixed_plan(conn),
    )
    db_session.add(report)
    await db_session.flush()

    async def fake_generate(provider: str, api_key: str, model: str, **kwargs: Any) -> ReportQuery:
        return _query("SELECT name FROM repositories")

    monkeypatch.setattr(reports_service, "generate_report_query", fake_generate)

    async def fake_docs_section(session: AsyncSession, **kwargs: Any) -> SectionResult:
        planned = kwargs["planned"]
        return SectionResult(
            title=planned.title,
            source="documents",
            sub_question=planned.sub_question,
            kind="text",
            prose="Відпустка триває 24 дні [1].",
        )

    monkeypatch.setattr(reports_service, "execute_documents_section", fake_docs_section)

    await reports_service.generate_report(db_session, report)

    assert report.status is ReportStatus.DONE
    assert report.pdf_data is not None and report.pdf_data[:5] == b"%PDF-"
    assert report.sections is not None and len(report.sections) == 2
    assert report.sections[0]["source"] == "database"
    assert report.sections[0]["generated_sql"] is not None
    assert report.sections[1]["source"] == "documents"
    # Legacy single-SQL columns stay empty on blended reports (schedules guard on them).
    assert report.generated_sql is None and report.chart_spec is None
    assert report.sections_data is not None
    # The plan carries a blank language (``_mixed_plan``); the stored payload must
    # hold the *resolved* language ("uk", inferred from the Cyrillic prompt), not
    # the raw blank value the plan still has on the row.
    assert report.sections_data["language"] == "uk"
    assert len(report.sections_data["sections"]) == len(report.sections)


async def test_a_report_with_nothing_in_any_section_is_not_produced(
    db_session: AsyncSession, monkeypatch: pytest.MonkeyPatch
) -> None:
    """A cover page, a heading and a sentence saying there was nothing to say.

    That is what a reader was handed (#185): a full generation run spent to
    produce a PDF with no content in it. If nothing matched anywhere, say so on
    the row instead of making them open a document to find out.
    """
    repo, conn = await _connected_repo(db_session)
    report = GeneratedReport(
        repository_id=repo.id,
        connection_id=conn.id,
        requested_by=None,
        prompt="огляд",
        plan=_mixed_plan(conn),
    )
    db_session.add(report)
    await db_session.flush()

    async def fake_generate(provider: str, api_key: str, model: str, **kwargs: Any) -> ReportQuery:
        return _query("SELECT name FROM repositories WHERE name = 'Tesla'")

    async def empty_docs(session: AsyncSession, **kwargs: Any) -> SectionResult:
        planned = kwargs["planned"]
        return SectionResult(
            title=planned.title,
            source="documents",
            sub_question=planned.sub_question,
            kind="text",
            prose=label("no_documents", "uk"),
            empty=True,
        )

    monkeypatch.setattr(reports_service, "generate_report_query", fake_generate)
    monkeypatch.setattr(reports_service, "execute_documents_section", empty_docs)

    await reports_service.generate_report(db_session, report)

    assert report.status is ReportStatus.FAILED
    # Distinguishable from a provider or timeout failure, which is the point:
    # the reader is told the repository held nothing, not that something broke.
    assert report.error is not None and "EmptyReportError" in report.error
    # Nothing to download. A refused report that still produced bytes would be
    # the empty document this card exists to stop.
    assert report.pdf_data is None
    assert report.pptx_data is None


async def test_one_section_with_data_still_produces_the_report(
    db_session: AsyncSession, monkeypatch: pytest.MonkeyPatch
) -> None:
    """The line this card must not cross. A three-section report with one empty
    section is still useful, and refusing it would throw away the two that
    worked."""
    repo, conn = await _connected_repo(db_session)
    report = GeneratedReport(
        repository_id=repo.id,
        connection_id=conn.id,
        requested_by=None,
        prompt="огляд",
        plan=_mixed_plan(conn),
    )
    db_session.add(report)
    await db_session.flush()

    async def fake_generate(provider: str, api_key: str, model: str, **kwargs: Any) -> ReportQuery:
        return _query("SELECT name FROM repositories")

    async def empty_docs(session: AsyncSession, **kwargs: Any) -> SectionResult:
        planned = kwargs["planned"]
        return SectionResult(
            title=planned.title,
            source="documents",
            sub_question=planned.sub_question,
            kind="text",
            prose=label("no_documents", "uk"),
            empty=True,
        )

    monkeypatch.setattr(reports_service, "generate_report_query", fake_generate)
    monkeypatch.setattr(reports_service, "run_validated_query", _matching_rows)
    monkeypatch.setattr(reports_service, "execute_documents_section", empty_docs)

    await reports_service.generate_report(db_session, report)

    assert report.status is ReportStatus.DONE
    assert report.pdf_data is not None and report.pdf_data[:5] == b"%PDF-"


async def test_a_nightly_run_that_matches_nothing_fails_its_row(
    db_session: AsyncSession, monkeypatch: pytest.MonkeyPatch
) -> None:
    """The frozen path hits this too, and cannot retry (#185).

    Owner's decision, recorded: the row appears in history as **failed with the
    reason**, rather than the night writing nothing at all. A silent night is
    quieter but makes "did the schedule fire?" unanswerable from the UI, and
    ``execute_frozen`` already creates the row before it can fail — writing
    nothing would mean deleting it again.
    """
    repo, connection = await _connected_repo(db_session)
    schedule = ReportSchedule(
        repository_id=repo.id,
        connection_id=connection.id,
        owner_id=None,
        prompt="Виторг",
        frozen_sql="SELECT 1",
        frozen_chart_spec=ChartSpec(chart_type="none", title="t").model_dump(),
        run_at_time=time(1, 0),
    )
    db_session.add(schedule)
    await db_session.commit()

    async def no_rows(**kwargs: Any) -> QueryResult:
        return QueryResult(columns=["fy"], rows=[])

    monkeypatch.setattr(reports_service, "run_validated_query", no_rows)

    report = await reports_service.execute_frozen(db_session, schedule=schedule)

    assert report.status is ReportStatus.FAILED
    assert report.error is not None and "EmptyReportError" in report.error
    assert report.pdf_data is None and report.pptx_data is None
    # The failed row is genuinely on the table: that is the whole point of
    # failing rather than writing nothing.
    stored = await db_session.get(GeneratedReport, report.id)
    assert stored is not None and stored.status is ReportStatus.FAILED


async def test_partial_success_when_one_section_fails(
    db_session: AsyncSession, monkeypatch: pytest.MonkeyPatch
) -> None:
    repo, conn = await _connected_repo(db_session)
    report = GeneratedReport(
        repository_id=repo.id,
        connection_id=conn.id,
        requested_by=None,
        prompt="огляд",
        plan=_mixed_plan(conn),
    )
    db_session.add(report)
    await db_session.flush()

    async def fake_generate(provider: str, api_key: str, model: str, **kwargs: Any) -> ReportQuery:
        return _query("SELECT name FROM repositories")

    monkeypatch.setattr(reports_service, "generate_report_query", fake_generate)

    async def exploding_docs_section(session: AsyncSession, **kwargs: Any) -> SectionResult:
        raise RuntimeError("provider exploded")

    monkeypatch.setattr(reports_service, "execute_documents_section", exploding_docs_section)

    await reports_service.generate_report(db_session, report)

    assert report.status is ReportStatus.DONE  # partial success
    assert report.sections is not None
    assert report.sections[1]["status"] == "error"
    assert "provider exploded" in (report.sections[1]["error"] or "")


async def test_all_sections_failed_fails_the_report(
    db_session: AsyncSession, monkeypatch: pytest.MonkeyPatch
) -> None:
    repo, conn = await _connected_repo(db_session)
    report = GeneratedReport(
        repository_id=repo.id,
        connection_id=conn.id,
        requested_by=None,
        prompt="огляд",
        plan={
            "title": "",
            "language": "",
            "sections": [{"title": "Docs", "source": "documents", "sub_question": "q"}],
        },
    )
    db_session.add(report)
    await db_session.flush()

    async def exploding_docs_section(session: AsyncSession, **kwargs: Any) -> SectionResult:
        raise RuntimeError("nope")

    monkeypatch.setattr(reports_service, "execute_documents_section", exploding_docs_section)

    await reports_service.generate_report(db_session, report)

    assert report.status is ReportStatus.FAILED
    assert report.error is not None and "nope" in report.error


async def test_docs_only_repo_generates_without_connection(
    db_session: AsyncSession, monkeypatch: pytest.MonkeyPatch
) -> None:
    repo = await _docs_only_repo(db_session)
    report = GeneratedReport(
        repository_id=repo.id,
        connection_id=None,
        requested_by=None,
        prompt="політика",
        plan={
            "title": "",
            "language": "",
            "sections": [{"title": "Політика", "source": "documents", "sub_question": "q"}],
        },
    )
    db_session.add(report)
    await db_session.flush()

    async def fake_docs_section(session: AsyncSession, **kwargs: Any) -> SectionResult:
        planned = kwargs["planned"]
        return SectionResult(
            title=planned.title,
            source="documents",
            sub_question=planned.sub_question,
            kind="text",
            prose="24 дні.",
        )

    monkeypatch.setattr(reports_service, "execute_documents_section", fake_docs_section)

    await reports_service.generate_report(db_session, report)

    assert report.status is ReportStatus.DONE
    assert report.connection_id is None
    assert report.pdf_data is not None


async def test_plan_report_stops_at_awaiting_approval(
    db_session: AsyncSession, monkeypatch: pytest.MonkeyPatch
) -> None:
    """Planning stores the plan and waits; it must not execute anything."""
    repo, conn = await _connected_repo(db_session)
    report = GeneratedReport(
        repository_id=repo.id, connection_id=conn.id, prompt="дохід", status=ReportStatus.PLANNING
    )
    db_session.add(report)
    await db_session.flush()

    async def fake_plan(provider: str, api_key: str, model: str, **kwargs: Any) -> ReportPlan:
        # The repository's databases are offered, each one named (#203).
        assert [d.id for d in kwargs["databases"]] == [str(conn.id)]
        assert kwargs["databases"][0].name == "Reporting"
        assert kwargs["has_documents"] is False
        return ReportPlan(
            title="Дохід",
            language="uk",
            sections=[
                PlannedSection(
                    title="Динаміка",
                    source="database",
                    databases=[str(conn.id)],
                    sub_question="дохід?",
                )
            ],
        )

    monkeypatch.setattr(reports_service, "generate_report_plan", fake_plan)

    async def boom(*args: Any, **kwargs: Any) -> None:
        raise AssertionError("generation must not run during planning")

    monkeypatch.setattr(reports_service, "build_report_pdf", boom)

    await reports_service.plan_report(db_session, report)

    assert report.status is ReportStatus.AWAITING_APPROVAL
    assert report.plan is not None
    assert report.plan["sections"][0]["title"] == "Динаміка"
    assert report.plan_rounds == 1
    assert report.pdf_data is None


async def test_plan_report_offers_only_the_capabilities_a_repo_has(
    db_session: AsyncSession, monkeypatch: pytest.MonkeyPatch
) -> None:
    """A docs-only repo must not be offered a DB the planner could route sections to."""
    repo = await _docs_only_repo(db_session)
    report = GeneratedReport(
        repository_id=repo.id, connection_id=None, prompt="політика", status=ReportStatus.PLANNING
    )
    db_session.add(report)
    await db_session.flush()

    async def fake_plan(provider: str, api_key: str, model: str, **kwargs: Any) -> ReportPlan:
        assert kwargs["databases"] == []  # no DB capability offered
        assert kwargs["has_documents"] is True
        return ReportPlan(
            sections=[PlannedSection(title="Політика", source="documents", sub_question="q")]
        )

    monkeypatch.setattr(reports_service, "generate_report_plan", fake_plan)

    await reports_service.plan_report(db_session, report)

    assert report.status is ReportStatus.AWAITING_APPROVAL
    assert report.plan is not None and report.plan["sections"][0]["source"] == "documents"


async def test_plan_report_records_failure_on_the_row(
    db_session: AsyncSession, monkeypatch: pytest.MonkeyPatch
) -> None:
    repo, conn = await _connected_repo(db_session)
    report = GeneratedReport(
        repository_id=repo.id, connection_id=conn.id, prompt="дохід", status=ReportStatus.PLANNING
    )
    db_session.add(report)
    await db_session.flush()

    async def always_bad(provider: str, api_key: str, model: str, **kwargs: Any) -> ReportPlan:
        raise ReportPlanParseError("Reply contained no JSON object.")

    monkeypatch.setattr(reports_service, "generate_report_plan", always_bad)

    await reports_service.plan_report(db_session, report)

    assert report.status is ReportStatus.FAILED
    assert report.error is not None
    assert report.plan is None


async def test_generate_report_uses_the_stored_plan_without_planning(
    db_session: AsyncSession, monkeypatch: pytest.MonkeyPatch
) -> None:
    """Generation reads the approved plan; it must never call the planner again."""
    repo, conn = await _connected_repo(db_session)
    report = GeneratedReport(
        repository_id=repo.id,
        connection_id=conn.id,
        prompt="назви репозиторіїв",
        status=ReportStatus.PENDING,
        plan={
            "title": "Репозиторії",
            "language": "uk",
            "sections": [
                {
                    "title": "Назви",
                    "source": "database",
                    "databases": [str(conn.id)],
                    "sub_question": "які назви?",
                }
            ],
        },
    )
    db_session.add(report)
    await db_session.flush()

    async def never(*args: Any, **kwargs: Any) -> ReportPlan:
        raise AssertionError("generation must not plan")

    monkeypatch.setattr(reports_service, "generate_report_plan", never)

    async def fake_query(provider: str, api_key: str, model: str, **kwargs: Any) -> ReportQuery:
        return ReportQuery(
            sql="SELECT name FROM repositories LIMIT 5",
            chart=ChartSpec(chart_type="none", title="Назви"),
            column_labels={"name": "Назва"},
        )

    monkeypatch.setattr(reports_service, "generate_report_query", fake_query)
    monkeypatch.setattr(reports_service, "run_validated_query", _matching_rows)

    await reports_service.generate_report(db_session, report)

    assert report.status is ReportStatus.DONE
    assert report.pdf_data is not None


async def test_generate_report_fails_without_a_plan(
    db_session: AsyncSession,
) -> None:
    repo, conn = await _connected_repo(db_session)
    report = GeneratedReport(
        repository_id=repo.id, connection_id=conn.id, prompt="p", status=ReportStatus.PENDING
    )
    db_session.add(report)
    await db_session.flush()

    await reports_service.generate_report(db_session, report)

    assert report.status is ReportStatus.FAILED
    assert report.error is not None and "plan" in report.error.lower()


async def test_replan_fills_blanks_and_keeps_user_edits(
    db_session: AsyncSession, monkeypatch: pytest.MonkeyPatch
) -> None:
    """The convergence lock, end to end: only the *provider* is stubbed.

    Stubbing ``generate_replan`` would prove nothing — the merge is the feature, so
    ``build_replan_prompt`` → ``parse_replan`` → ``merge_draft`` must all really
    run. The reply below is maximally hostile: it rewrites the report title, the
    language, both section titles, a sub-question the user wrote, and flips both
    sources. Every one of those is discarded; only the blank sub-question takes the
    model's text. The prompt assertions pin the other half of the property — the
    model is asked about the *submitted draft*, never the plan sitting on the row.
    """
    repo, conn = await _connected_repo(db_session)
    report = GeneratedReport(
        repository_id=repo.id,
        connection_id=conn.id,
        prompt="дохід",
        status=ReportStatus.PLANNING,
        plan={
            "title": "Заголовок минулого раунду",
            "language": "uk",
            "sections": [
                {
                    "title": "Стара динаміка",
                    "source": "database",
                    "databases": [str(conn.id)],
                    "sub_question": "старе питання",
                },
                {
                    "title": "Стара маржа",
                    "source": "database",
                    "databases": [str(conn.id)],
                    "sub_question": "старе питання 2",
                },
            ],
        },
        plan_rounds=1,
    )
    db_session.add(report)
    await db_session.flush()

    draft = ReportPlanDraft(
        title="Мій заголовок",
        language="uk",
        sections=[
            DraftSection(
                title="Динаміка",
                source="database",
                databases=[str(conn.id)],
                sub_question="дохід за FY25?",
            ),
            DraftSection(
                title="Маржа", source="database", databases=[str(conn.id)], sub_question=""
            ),
        ],
    )

    prompts: list[str] = []

    async def fake_generate_text(
        provider: str, api_key: str, model: str, *, prompt: str, base_url: str | None = None
    ) -> str:
        prompts.append(prompt)
        return json.dumps(
            {
                "title": "Заголовок від моделі",
                "language": "en",
                "sections": [
                    {
                        "title": "Динаміка від моделі",
                        "source": "documents",
                        "sub_question": "Модель переписала це питання",
                    },
                    {
                        "title": "Маржа від моделі",
                        "source": "documents",
                        "sub_question": "Яка маржа у FY25?",
                    },
                ],
            },
            ensure_ascii=False,
        )

    monkeypatch.setattr(report_planner, "generate_text", fake_generate_text)

    await reports_service.replan_report(db_session, report, draft)

    assert report.status is ReportStatus.AWAITING_APPROVAL
    assert report.plan == {
        "title": "Мій заголовок",
        "language": "uk",
        "sections": [
            {
                "title": "Динаміка",
                "source": "database",
                "databases": [str(conn.id)],
                # #204's output contract: empty, because this section reads one
                # database and so has nothing to concatenate its rows with.
                "columns": [],
                "sub_question": "дохід за FY25?",
            },
            # The one genuinely blank field, and the only one the model filled.
            {
                "title": "Маржа",
                "source": "database",
                "databases": [str(conn.id)],
                "columns": [],
                "sub_question": "Яка маржа у FY25?",
            },
        ],
    }
    assert report.plan_rounds == 2
    assert "дохід за FY25?" in prompts[0]  # the user's draft went to the model…
    assert "Стара динаміка" not in prompts[0]  # …not the plan already on the row


async def test_replan_without_blanks_makes_no_provider_call(
    db_session: AsyncSession, monkeypatch: pytest.MonkeyPatch
) -> None:
    """A complete draft is stored as-is: there is nothing for the model to do."""
    repo, conn = await _connected_repo(db_session)
    report = GeneratedReport(
        repository_id=repo.id,
        connection_id=conn.id,
        prompt="дохід",
        status=ReportStatus.PLANNING,
        plan_rounds=1,
    )
    db_session.add(report)
    await db_session.flush()

    async def boom(*args: Any, **kwargs: Any) -> ReportPlan:
        raise AssertionError("a complete draft must not reach the provider")

    monkeypatch.setattr(reports_service, "generate_replan", boom)

    draft = ReportPlanDraft(
        title="Повний план",
        language="uk",
        # Routed, or it would not be complete: since #203 a database section with
        # no database is a blank like any other, and this draft would buy a round.
        sections=[
            DraftSection(
                title="Динаміка",
                source="database",
                databases=[str(conn.id)],
                sub_question="дохід?",
            )
        ],
    )
    await reports_service.replan_report(db_session, report, draft)

    assert report.status is ReportStatus.AWAITING_APPROVAL
    assert report.plan is not None and report.plan["title"] == "Повний план"
    assert report.plan_rounds == 1  # unchanged: no round was paid for


async def test_replan_records_failure_on_the_row(
    db_session: AsyncSession, monkeypatch: pytest.MonkeyPatch
) -> None:
    """With no plan to fall back on there is nothing to preserve: FAILED, as before."""
    repo, conn = await _connected_repo(db_session)
    report = GeneratedReport(
        repository_id=repo.id, connection_id=conn.id, prompt="дохід", status=ReportStatus.PLANNING
    )
    db_session.add(report)
    await db_session.flush()

    async def always_bad(*args: Any, **kwargs: Any) -> ReportPlan:
        raise ReportPlanParseError("Returned 3 sections, but the plan has 2.")

    monkeypatch.setattr(reports_service, "generate_replan", always_bad)

    draft = ReportPlanDraft(
        title="T",
        language="uk",
        sections=[DraftSection(title="A", source="database", sub_question="")],
    )
    await reports_service.replan_report(db_session, report, draft)

    assert report.status is ReportStatus.FAILED
    assert report.error is not None


async def test_replan_failure_leaves_the_previous_plan_approvable(
    db_session: AsyncSession, monkeypatch: pytest.MonkeyPatch
) -> None:
    """A failed round must not destroy a plan the user could still approve.

    ``FAILED`` is terminal — nothing moves a report out of it and both plan
    endpoints require ``AWAITING_APPROVAL`` — while the user's edited draft was
    never persisted. So a provider that fails every attempt (here by breaking the
    section-count contract three times, exactly what deleting a section provokes)
    would otherwise permanently kill a perfectly approvable report.
    """
    repo, conn = await _connected_repo(db_session)
    previous_plan = {
        "title": "Затверджуваний план",
        "language": "uk",
        "sections": [{"title": "Динаміка", "source": "database", "sub_question": "дохід?"}],
    }
    report = GeneratedReport(
        repository_id=repo.id,
        connection_id=conn.id,
        prompt="дохід",
        status=ReportStatus.PLANNING,
        plan=dict(previous_plan),
        plan_rounds=1,
    )
    db_session.add(report)
    await db_session.flush()

    async def always_down(
        provider: str, api_key: str, model: str, *, prompt: str, base_url: str | None = None
    ) -> str:
        # A reply with one section too many: rejected by ``merge_draft`` on every
        # attempt, so the repair loop exhausts its budget.
        return json.dumps(
            {
                "title": "t",
                "language": "uk",
                "sections": [
                    {"title": "A", "source": "database", "sub_question": "q"},
                    {"title": "B", "source": "database", "sub_question": "q"},
                ],
            }
        )

    monkeypatch.setattr(report_planner, "generate_text", always_down)

    draft = ReportPlanDraft(
        title="Мій заголовок",
        language="uk",
        sections=[DraftSection(title="Динаміка", source="database", sub_question="")],
    )
    await reports_service.replan_report(db_session, report, draft)
    await db_session.refresh(report)

    assert report.status is ReportStatus.AWAITING_APPROVAL
    assert report.error is not None and "sections" in report.error
    assert report.plan == previous_plan  # the round failed; the plan survived
    assert report.plan_rounds == 1  # no round was completed


async def test_a_query_matching_nothing_is_retried_once_with_its_literals(
    db_session: AsyncSession, monkeypatch: pytest.MonkeyPatch
) -> None:
    """Eleven of a hundred campaign reports were empty because a literal was guessed.

    `filings.company` holds 'Apple Inc.', not 'Apple'; a company stops reporting
    an XBRL tag and starts using another. Zero rows is not an error, so nothing
    upstream noticed and the report was delivered as a success containing the
    single line "Rows: 0".
    """
    repo, conn = await _connected_repo(db_session)
    report = GeneratedReport(
        repository_id=repo.id,
        connection_id=conn.id,
        requested_by=None,
        prompt="names",
        plan=_single_db_plan("names", conn),
    )
    db_session.add(report)
    await db_session.flush()

    attempts: list[list[str]] = []
    replies = iter(
        [
            _query("SELECT name FROM repositories WHERE name = 'Apple'"),
            _query("SELECT name FROM repositories"),
        ]
    )

    async def fake_generate(provider: str, api_key: str, model: str, **kwargs: Any) -> ReportQuery:
        attempts.append(list(kwargs["feedback"]))
        return next(replies)

    monkeypatch.setattr(reports_service, "generate_report_query", fake_generate)
    monkeypatch.setattr(reports_service, "run_validated_query", _matching_rows)
    monkeypatch.setattr(
        reports_service.provider_service,  # type: ignore[attr-defined]
        "get_provider_key",
        _fake_key,
    )

    await reports_service.generate_report(db_session, report)

    assert len(attempts) == 2, attempts
    complaint = attempts[1][0]
    assert "NO ROWS" in complaint
    assert "'Apple'" in complaint  # the literal it actually filtered on
    assert report.status is ReportStatus.DONE
    assert report.sections is not None
    assert "WHERE" not in (report.sections[0]["generated_sql"] or "")  # the broader query won


async def _grounded_docs(session: AsyncSession, **kwargs: Any) -> SectionResult:
    """A documents section that DID find something.

    Used by the two empty-result tests below. Their subject is the database
    section's retry behaviour, and since #185 a report whose sections are *all*
    empty is refused before the section audit is ever written — so they pair the
    empty section with a full one. That keeps the report alive, keeps
    ``report.sections`` assertable, and leaves each test failing for its own
    reason rather than for #185's.
    """
    planned = kwargs["planned"]
    return SectionResult(
        title=planned.title,
        source="documents",
        sub_question=planned.sub_question,
        kind="text",
        prose="Відпустка триває 24 дні [1].",
    )


async def test_a_genuinely_empty_answer_is_kept_rather_than_failed(
    db_session: AsyncSession, monkeypatch: pytest.MonkeyPatch
) -> None:
    """Sometimes nothing IS the answer — Tesla is not in this data.

    The retry must not turn an honest empty result into a failed section, and
    must not keep spending attempts once it has been asked twice.
    """
    repo, conn = await _connected_repo(db_session)
    report = GeneratedReport(
        repository_id=repo.id,
        connection_id=conn.id,
        requested_by=None,
        prompt="names",
        plan=_mixed_plan(conn),
    )
    db_session.add(report)
    await db_session.flush()

    calls = 0

    async def always_empty(provider: str, api_key: str, model: str, **kwargs: Any) -> ReportQuery:
        nonlocal calls
        calls += 1
        return _query("SELECT name FROM repositories WHERE name = 'Tesla'")

    monkeypatch.setattr(reports_service, "generate_report_query", always_empty)
    monkeypatch.setattr(reports_service, "execute_documents_section", _grounded_docs)
    monkeypatch.setattr(
        reports_service.provider_service,  # type: ignore[attr-defined]
        "get_provider_key",
        _fake_key,
    )

    await reports_service.generate_report(db_session, report)

    assert calls == 2  # asked once more, then accepted — not all three attempts
    assert report.status is ReportStatus.DONE
    assert report.sections is not None
    assert report.sections[0]["status"] == "done"  # honest empty, not an error


async def test_an_empty_answer_survives_later_attempts_being_rejected(
    db_session: AsyncSession, monkeypatch: pytest.MonkeyPatch
) -> None:
    """The empty result is held onto, not just re-derived by a later attempt.

    First query runs and matches nothing; both retries are rejected by the
    guardrails. Without keeping the first result the section becomes "could not
    produce a valid query" — an error, when a valid query had in fact run and the
    honest answer was "nothing matched".
    """
    repo, conn = await _connected_repo(db_session)
    report = GeneratedReport(
        repository_id=repo.id,
        connection_id=conn.id,
        requested_by=None,
        prompt="names",
        plan=_mixed_plan(conn),
    )
    db_session.add(report)
    await db_session.flush()

    replies = iter(
        [
            _query("SELECT name FROM repositories WHERE name = 'Tesla'"),
            _query("SELECT secret FROM vault_x"),
            _query("SELECT secret FROM vault_y"),
        ]
    )

    async def fake_generate(provider: str, api_key: str, model: str, **kwargs: Any) -> ReportQuery:
        return next(replies)

    monkeypatch.setattr(reports_service, "generate_report_query", fake_generate)
    monkeypatch.setattr(reports_service, "execute_documents_section", _grounded_docs)
    monkeypatch.setattr(
        reports_service.provider_service,  # type: ignore[attr-defined]
        "get_provider_key",
        _fake_key,
    )

    await reports_service.generate_report(db_session, report)

    assert report.status is ReportStatus.DONE
    assert report.sections is not None
    assert report.sections[0]["status"] == "done", report.sections[0]
    # The point of this test: the *held* result is delivered rather than
    # re-derived, so the SQL that actually ran is the one on the row.
    assert "Tesla" in (report.sections[0]["generated_sql"] or "")


async def test_ordinary_return_types_do_not_fail_the_whole_report(
    db_session: AsyncSession, monkeypatch: pytest.MonkeyPatch
) -> None:
    """``SELECT id, created_at, total FROM orders`` on a UUID-keyed table.

    Every value here is an ordinary asyncpg return type. Persisting them runs
    inside ``generate_report``'s try block, so a codec that cannot encode one of
    them does not merely skip the payload — it marks a report FAILED that would
    have rendered fine the day before the column existed. The NaN is the nastier
    half: ``json.dumps`` emits it happily and Postgres rejects it at COMMIT.
    """
    repo, conn = await _connected_repo(db_session)
    report = GeneratedReport(
        repository_id=repo.id,
        connection_id=conn.id,
        requested_by=None,
        prompt="orders",
        plan=_single_db_plan("orders", conn),
    )
    db_session.add(report)
    await db_session.flush()

    async def fake_generate(provider: str, api_key: str, model: str, **kwargs: Any) -> ReportQuery:
        return _query("SELECT name FROM repositories")

    async def fake_query(**kwargs: Any) -> QueryResult:
        return QueryResult(
            # ``score``, not ``ratio``: a NaN in a percent-named column trips a
            # pre-existing ``InvalidOperation`` in ``report_format`` that has
            # nothing to do with persistence (see the fix report).
            columns=["id", "blob", "at", "window", "amounts", "meta", "score"],
            rows=[
                (
                    uuid.uuid4(),
                    b"\x00\xff",
                    time(9, 30),
                    timedelta(hours=3),
                    [Decimal("1.5"), Decimal("2.5")],  # numeric[]
                    {"total": Decimal("3.5")},  # jsonb
                    float("nan"),
                )
            ],
        )

    monkeypatch.setattr(reports_service, "generate_report_query", fake_generate)
    monkeypatch.setattr(reports_service, "run_validated_query", fake_query)
    monkeypatch.setattr(
        reports_service.provider_service,  # type: ignore[attr-defined]
        "get_provider_key",
        _fake_key,
    )

    await reports_service.generate_report(db_session, report)

    assert report.status is ReportStatus.DONE, report.error
    assert report.sections_data is not None


def test_a_payload_that_cannot_be_serialized_costs_only_the_rebuild(
    report_row: GeneratedReport, monkeypatch: pytest.MonkeyPatch
) -> None:
    """Belt and braces: persistence must never be able to fail a report.

    If some future value defeats the codec anyway, the honest outcome is "this
    report is not re-renderable" — the same degradation the size budget already
    takes — not a FAILED report with no artifacts.
    """

    class Unserializable:
        def __repr__(self) -> str:
            return "<obj>"

    def hostile(*args: Any, **kwargs: Any) -> Any:
        raise TypeError("Object of type Unserializable is not JSON serializable")

    monkeypatch.setattr(reports_service, "dump_sections", hostile)
    _store_sections(report_row, title="Звіт", language="uk", sections=[])
    assert report_row.sections_data is None


async def test_frozen_render_stores_the_payload_and_the_audit(
    db_session: AsyncSession, monkeypatch: pytest.MonkeyPatch
) -> None:
    """The frozen-schedule path (``execute_frozen`` → ``_run_and_render``) wrote no
    ``sections`` audit at all before this change — closed here, or a scheduled
    report could never be re-rendered.
    """
    _, connection = await _connected_repo(db_session)

    async def fake_query(**kwargs: Any) -> QueryResult:
        return QueryResult(columns=["fy", "revenue"], rows=[(2024, Decimal("391.0"))])

    monkeypatch.setattr(reports_service, "run_validated_query", fake_query)

    report = GeneratedReport(prompt="Виторг за 2024", repository_id=connection.repository_id)
    await reports_service._run_and_render(
        report,
        connection,
        "SELECT 1",
        ChartSpec(chart_type="bar", x_column="fy", y_column="revenue"),
        style=resolve_style(None),
    )

    assert report.status is ReportStatus.DONE
    assert report.sections is not None and len(report.sections) == 1
    assert report.sections_data is not None
    assert len(report.sections_data["sections"]) == 1


async def test_a_frozen_reports_chart_is_redrawn_byte_for_byte(
    db_session: AsyncSession, monkeypatch: pytest.MonkeyPatch
) -> None:
    """The branch's central promise, on the path that nearly lost it.

    ``render_chart`` suppresses a chart title that merely echoes the headings
    printed above the image (``keeps_title``). Generation and re-render must
    therefore agree on which headings they pass, or a frozen report with a chart
    title — the normal case — comes back with a differently laid-out figure. The
    only way to be sure they agree is for both to go through ``_with_chart``.
    """
    _, connection = await _connected_repo(db_session)

    async def fake_query(**kwargs: Any) -> QueryResult:
        return QueryResult(
            columns=["fy", "revenue"], rows=[(2024, Decimal("391.0")), (2023, Decimal("383.3"))]
        )

    monkeypatch.setattr(reports_service, "run_validated_query", fake_query)

    rendered: list[SectionResult] = []

    def capture(
        *, title: str, sections: list[SectionResult], language: str, **_style: object
    ) -> bytes:
        rendered.extend(sections)
        return b"%PDF-captured"

    monkeypatch.setattr(reports_service, "build_report_pdf", capture)

    report = GeneratedReport(prompt="Виторг за рік", repository_id=connection.repository_id)
    await reports_service._run_and_render(
        report,
        connection,
        "SELECT 1",
        ChartSpec(
            chart_type="bar", x_column="fy", y_column="revenue", title="Динаміка виторгу за рік"
        ),
        style=resolve_style(None),
    )

    assert report.sections_data is not None
    stored_title, stored_language, stored = load_sections(report.sections_data)
    redrawn = reports_service._with_chart(
        stored[0],
        language=stored_language,
        report_title=stored_title,
        style=resolve_style(None),
    )
    assert rendered[0].chart_png is not None
    assert redrawn.chart_png == rendered[0].chart_png


async def test_a_frozen_run_that_fails_while_rendering_keeps_no_payload(
    db_session: AsyncSession, monkeypatch: pytest.MonkeyPatch
) -> None:
    """A FAILED report must never advertise itself as re-renderable.

    ``_store_sections`` runs before the renderers, so a renderer that raises left
    the payload written on a row that then committed as FAILED with no artifacts:
    ``can_rerender`` said true, and ``/rerender`` would have manufactured a PDF
    for a report its owner was told had failed.
    """
    repo, connection = await _connected_repo(db_session)
    schedule = ReportSchedule(
        repository_id=repo.id,
        connection_id=connection.id,
        owner_id=None,
        prompt="Виторг",
        frozen_sql="SELECT 1",
        frozen_chart_spec=ChartSpec(chart_type="none", title="t").model_dump(),
        run_at_time=time(1, 0),
    )
    db_session.add(schedule)
    await db_session.commit()
    schedule_id = schedule.id

    async def fake_query(**kwargs: Any) -> QueryResult:
        return QueryResult(columns=["fy"], rows=[(2024,)])

    def boom(**kwargs: Any) -> bytes:
        raise RuntimeError("the renderer exploded")

    monkeypatch.setattr(reports_service, "run_validated_query", fake_query)
    monkeypatch.setattr(reports_service, "build_report_pdf", boom)

    report = await reports_service.execute_frozen(db_session, schedule=schedule)

    assert report.status is ReportStatus.FAILED
    assert report.sections_data is None
    # And the FAILED row is genuinely on the table — the rollback must not take
    # the report with it, or the schedule's owner would see no report at all.
    stored = (
        await db_session.execute(
            sa.select(GeneratedReport).where(GeneratedReport.schedule_id == schedule_id)
        )
    ).scalar_one()
    assert stored.status is ReportStatus.FAILED
    assert stored.sections_data is None
    assert stored.error is not None and "exploded" in stored.error


async def test_rerender_refuses_a_report_that_is_not_done(
    db_session: AsyncSession, stored_report: GeneratedReport
) -> None:
    # Belt to the braces above: whatever put section data on a report that did not
    # finish, rebuilding it would hand its owner a PDF for a failed report.
    stored_report.status = ReportStatus.FAILED
    await db_session.commit()
    with pytest.raises(ReportNotRenderable):
        await rerender_report(stored_report.id, session_factory=_fixed_factory(db_session))


@pytest.fixture
async def stored_report(
    db_session: AsyncSession, one_data_section: SectionResult
) -> GeneratedReport:
    """A DONE report whose sections_data payload alone can rebuild both artifacts.

    No plan, no connection: a report re-rendered from storage must not need
    either — the payload envelope carries the title and language on its own
    (see ``report_persist.dump_sections``).

    The prompt is deliberately English while the envelope says ``uk``: the
    fallback ``resolve_language`` would reach for is the prompt's script, so a
    re-render that consulted anything but the envelope would come back in the
    wrong language, and the test below would catch it.
    """
    repo = Repository(name="Vault", llm_provider="gemini", llm_model="gem-1")
    db_session.add(repo)
    await db_session.flush()
    report = GeneratedReport(
        repository_id=repo.id,
        prompt="Revenue for 2024",
        status=ReportStatus.DONE,
    )
    _store_sections(report, title="Звіт", language="uk", sections=[one_data_section])
    db_session.add(report)
    await db_session.commit()
    return report


async def test_rerender_touches_neither_the_provider_nor_the_database(
    db_session: AsyncSession,
    monkeypatch: pytest.MonkeyPatch,
    stored_report: GeneratedReport,
) -> None:
    """The whole guarantee of Task 5: stub the provider *and* the reporting-database
    query so either one being called fails the test. Asserting only "the PDF
    changed" would pass even if re-render quietly re-ran everything.
    """

    def explode(*args: Any, **kwargs: Any) -> Any:
        raise AssertionError("re-render must not call out")

    monkeypatch.setattr("contextvault.services.reports.run_validated_query", explode)
    monkeypatch.setattr("contextvault.services.report_llm.generate_text", explode)

    await rerender_report(stored_report.id, session_factory=_fixed_factory(db_session))

    await db_session.refresh(stored_report)
    assert stored_report.pdf_data is not None
    assert stored_report.pptx_data is not None
    assert stored_report.status == ReportStatus.DONE


async def test_rerender_refuses_a_report_with_no_stored_data(
    db_session: AsyncSession, stored_report: GeneratedReport
) -> None:
    stored_report.sections_data = None
    await db_session.commit()
    with pytest.raises(ReportNotRenderable):
        await rerender_report(stored_report.id, session_factory=_fixed_factory(db_session))


def _capture_pdf(monkeypatch: pytest.MonkeyPatch) -> dict[str, Any]:
    """Record what ``rerender_report`` hands the PDF renderer.

    The renderer's arguments are the only place the envelope's decisions are
    observable: the finished PDF cannot be asserted on usefully, because the
    renderer zlib-compresses its content streams — a test looking for a Cyrillic
    byte in the output finds one in an all-English PDF too, and so can never
    fail. (True of fpdf2 before this and of WeasyPrint after it.)
    """
    captured: dict[str, Any] = {}

    def capture(
        *, title: str, sections: list[SectionResult], language: str, **rest: object
    ) -> bytes:
        captured.update(title=title, language=language, sections=sections, **rest)
        return b"%PDF-captured"

    monkeypatch.setattr(reports_service, "build_report_pdf", capture)
    return captured


async def test_rerender_takes_its_language_and_title_from_the_stored_envelope(
    db_session: AsyncSession, monkeypatch: pytest.MonkeyPatch, stored_report: GeneratedReport
) -> None:
    # A Ukrainian report must not come back in English: the language rides in the
    # payload envelope, not in the plan, because frozen reports have no plan. This
    # report's own prompt is English, so anything but the envelope gives "en".
    captured = _capture_pdf(monkeypatch)
    envelope = stored_report.sections_data
    assert envelope is not None

    await rerender_report(stored_report.id, session_factory=_fixed_factory(db_session))

    assert captured["language"] == "uk"
    assert captured["title"] == envelope["title"] == "Звіт"


async def test_a_rerendered_data_section_still_carries_its_chart(
    db_session: AsyncSession, monkeypatch: pytest.MonkeyPatch, stored_report: GeneratedReport
) -> None:
    # ``chart_spec`` is what decides whether ``_with_chart`` draws anything at all.
    # Without this, dropping it from the stored payload would silently strip every
    # chart from every re-rendered report and leave the suite green.
    captured = _capture_pdf(monkeypatch)

    await rerender_report(stored_report.id, session_factory=_fixed_factory(db_session))

    assert captured["sections"][0].chart_png is not None


async def test_a_rerender_reprints_the_original_cover_date_not_todays(
    db_session: AsyncSession, monkeypatch: pytest.MonkeyPatch, stored_report: GeneratedReport
) -> None:
    """A rebuild reproduces the report; a restamped cover would quietly not.

    ``_cover_date`` falls back to today when the row has no ``created_at`` yet,
    which is right while a report is being created and wrong for every rebuild
    afterwards. This is the assertion that separates the two cases.
    """
    stored_report.created_at = datetime(2026, 3, 1, tzinfo=UTC)
    await db_session.commit()
    captured = _capture_pdf(monkeypatch)

    await rerender_report(stored_report.id, session_factory=_fixed_factory(db_session))

    assert captured["report_date"] == date(2026, 3, 1)


async def test_generation_and_rerender_resolve_the_same_template(
    db_session: AsyncSession, monkeypatch: pytest.MonkeyPatch, stored_report: GeneratedReport
) -> None:
    """Two independent resolutions can disagree, and nothing else would notice.

    3a left this asserting the default on both sides, which stayed true no
    matter which side learned to read the column. Now each side is pinned to a
    *non-default* template so only genuine agreement passes.
    """
    repo = await db_session.get(Repository, stored_report.repository_id)
    assert repo is not None
    repo.report_template = design_payload("mahistral")
    stored_report.design = None
    await db_session.commit()

    generated = await reports_service._bind_style(db_session, stored_report)
    await db_session.commit()
    captured = _capture_pdf(monkeypatch)
    await rerender_report(stored_report.id, session_factory=_fixed_factory(db_session))

    assert generated is TEMPLATES["mahistral"]
    assert captured["style"] is generated


async def test_a_new_report_inherits_the_repository_template(
    db_session: AsyncSession, monkeypatch: pytest.MonkeyPatch, stored_report: GeneratedReport
) -> None:
    """Generation resolves the repository default and records it on the row."""
    repo = await db_session.get(Repository, stored_report.repository_id)
    assert repo is not None
    repo.report_template = design_payload("chasopys")
    stored_report.design = None
    await db_session.commit()

    style = await reports_service._bind_style(db_session, stored_report)

    assert style is TEMPLATES["chasopys"]
    assert stored_report.design == {"template": "chasopys"}


async def test_a_design_already_on_the_row_beats_the_repository_default(
    db_session: AsyncSession, stored_report: GeneratedReport
) -> None:
    """Storing a design before generation finishes pre-selects the template."""
    repo = await db_session.get(Repository, stored_report.repository_id)
    assert repo is not None
    repo.report_template = design_payload("chasopys")
    stored_report.design = design_payload("prezydia")
    await db_session.commit()

    style = await reports_service._bind_style(db_session, stored_report)

    assert style is TEMPLATES["prezydia"]


async def test_binding_applies_the_tone_the_reader_chose_not_only_the_template(
    db_session: AsyncSession, stored_report: GeneratedReport
) -> None:
    """**Generation must honour the whole stored design, not just its template.**

    The design step lets a reader set typography, page and brand tokens *before*
    the report is generated, and promises the first artifact already wears them.
    ``_bind_style`` resolved only ``design.template`` — so scale, margins and
    accent were silently discarded at generation and then appeared out of nowhere
    on the first rerender, which is the one path that called ``style_for``.

    The neighbouring template test asserts ``style is TEMPLATES[...]`` and still
    holds: ``apply_tokens`` returns the template unchanged when nothing is set,
    which is what keeps every report written before this rendering identically.
    That identity is also why that test could never have caught this.
    """
    # 1.25, not 1.4: the scale clamps at 1.30, and a test that straddles the clamp
    # would be measuring the clamp rather than whether tokens are applied at all.
    stored_report.design = {**design_payload("prezydia"), "typography": {"scale": 1.25}}
    await db_session.commit()

    style = await reports_service._bind_style(db_session, stored_report)

    assert style.body_pt == pytest.approx(TEMPLATES["prezydia"].body_pt * 1.25)
    assert style.h1_pt == pytest.approx(TEMPLATES["prezydia"].h1_pt * 1.25)


async def test_binding_records_what_resolved_not_what_was_asked_for(
    db_session: AsyncSession, stored_report: GeneratedReport
) -> None:
    """A slug that fell back must not stay on the row claiming otherwise.

    Without this the row says ``ghost`` while the artifacts say Кварта, and the
    disagreement is invisible until someone reads the column.
    """
    stored_report.design = {"template": "ghost"}
    await db_session.commit()

    style = await reports_service._bind_style(db_session, stored_report)

    assert style is resolve_style(None)
    assert stored_report.design == {"template": resolve_style(None).key}


async def test_a_rerender_ignores_a_repository_default_that_moved(
    db_session: AsyncSession, monkeypatch: pytest.MonkeyPatch, stored_report: GeneratedReport
) -> None:
    """**The load-bearing test of this step.**

    A finished report wears what it was generated in, forever. If re-render
    resolved the repository instead of the row, changing a repository setting
    would silently restyle every report ever generated from it — the one thing
    ``rerender`` promises not to do (``reports.py`` says so in a comment above
    the call site).
    """
    stored_report.design = design_payload("prezydia")
    repo = await db_session.get(Repository, stored_report.repository_id)
    assert repo is not None
    repo.report_template = design_payload("minimum")  # the default moves...
    await db_session.commit()
    captured = _capture_pdf(monkeypatch)

    await rerender_report(stored_report.id, session_factory=_fixed_factory(db_session))

    assert captured["style"] is TEMPLATES["prezydia"]  # ...and the report does not


async def test_re_binding_a_style_keeps_a_stored_block_list(
    db_session: AsyncSession, stored_report: GeneratedReport
) -> None:
    """A re-plan runs _bind_style again on an existing row. Replacing the whole
    payload there would silently discard a reader's edits.

    The starting design deliberately omits ``template`` — the repository has no
    default either, so ``_bind_style`` must resolve Кварта and *write it back*.
    Asserting on ``stored_report.design`` in memory would pass even if that
    write never reached the row: SQLAlchemy does not detect an in-place
    mutation of a JSONB column, so an in-place ``.update()`` regression leaves
    the attribute looking right in Python while the database keeps the old
    value forever. A raw SQL read after commit is the only thing that would
    catch it — it bypasses the ORM identity map entirely, so it can only see
    what was actually flushed.
    """
    stored_report.design = {"blocks_version": 1, "blocks": [{"kind": "cover"}]}
    await db_session.commit()

    await reports_service._bind_style(db_session, stored_report)
    await db_session.commit()

    design = (
        await db_session.execute(
            sa.text("SELECT design FROM generated_reports WHERE id = :id"),
            {"id": stored_report.id},
        )
    ).scalar_one()
    assert design["blocks"] == [{"kind": "cover"}]
    assert design["template"] == "kvarta"


async def test_a_rerender_of_a_legacy_row_is_the_default_template(
    db_session: AsyncSession, monkeypatch: pytest.MonkeyPatch, stored_report: GeneratedReport
) -> None:
    """Every report generated before this column renders exactly as it did."""
    stored_report.design = None
    await db_session.commit()
    captured = _capture_pdf(monkeypatch)

    await rerender_report(stored_report.id, session_factory=_fixed_factory(db_session))

    assert captured["style"] is resolve_style(None)


async def test_generation_renders_in_the_repository_template_and_records_it(
    db_session: AsyncSession, monkeypatch: pytest.MonkeyPatch
) -> None:
    """``generate_report`` actually binds — end to end, not through ``_bind_style``.

    Every other test in this group calls ``_bind_style`` directly, which leaves
    the seam between it and the orchestrator uncovered: deleting the call from
    ``generate_report`` and restoring ``resolve_style(None)`` would keep them all
    green while every report silently rendered in Кварта again.
    """
    repo, conn = await _connected_repo(db_session)
    repo.report_template = design_payload("chasopys")
    report = GeneratedReport(
        repository_id=repo.id,
        connection_id=conn.id,
        prompt="назви репозиторіїв",
        status=ReportStatus.PENDING,
        plan=_single_db_plan("які назви?", conn),
    )
    db_session.add(report)
    await db_session.flush()

    async def fake_query(provider: str, api_key: str, model: str, **kwargs: Any) -> ReportQuery:
        return ReportQuery(
            sql="SELECT name FROM repositories LIMIT 5",
            chart=ChartSpec(chart_type="none", title="Назви"),
        )

    monkeypatch.setattr(reports_service, "generate_report_query", fake_query)
    monkeypatch.setattr(reports_service, "run_validated_query", _matching_rows)
    captured = _capture_pdf(monkeypatch)

    await reports_service.generate_report(db_session, report)

    assert report.status is ReportStatus.DONE
    assert captured["style"] is TEMPLATES["chasopys"]
    assert report.design == {"template": "chasopys"}


async def test_a_frozen_run_renders_in_the_repository_template_and_records_it(
    db_session: AsyncSession, monkeypatch: pytest.MonkeyPatch
) -> None:
    """The same seam on the frozen-schedule path, which has no plan and no LLM call.

    ``_run_and_render`` used to resolve its own style; now it is handed one, so
    only driving ``execute_frozen`` proves the hand-off is wired up.
    """
    repo, connection = await _connected_repo(db_session)
    repo.report_template = design_payload("minimum")
    schedule = ReportSchedule(
        repository_id=repo.id,
        connection_id=connection.id,
        owner_id=None,
        prompt="Виторг",
        frozen_sql="SELECT 1",
        frozen_chart_spec=ChartSpec(chart_type="none", title="t").model_dump(),
        run_at_time=time(1, 0),
    )
    db_session.add(schedule)
    await db_session.commit()

    async def fake_query(**kwargs: Any) -> QueryResult:
        return QueryResult(columns=["fy"], rows=[(2024,)])

    monkeypatch.setattr(reports_service, "run_validated_query", fake_query)
    captured = _capture_pdf(monkeypatch)

    report = await reports_service.execute_frozen(db_session, schedule=schedule)

    assert report.status is ReportStatus.DONE
    assert captured["style"] is TEMPLATES["minimum"]
    assert report.design == {"template": "minimum"}


def _design_with_chart_type(chart_type: str) -> dict[str, Any]:
    return {
        "template": "kvarta",
        "blocks_version": BLOCKS_VERSION,
        "blocks": [
            {"kind": "cover"},
            {"kind": "heading", "section": 0},
            {"kind": "viz", "section": 0, "variant": "chart", "chart_type": chart_type},
            {"kind": "viz", "section": 0, "variant": "table"},
        ],
    }


async def test_a_stored_chart_type_override_reaches_the_drawn_figure(
    db_session: AsyncSession, stored_report: GeneratedReport
) -> None:
    """The override is applied where the figure is drawn (``_with_chart``), so it
    must actually change the bytes matplotlib produces on a real render — not
    merely round-trip through storage. Neither ``build_report_pdf`` nor
    ``rerender_report`` is stubbed here for that reason.
    """
    stored_report.design = _design_with_chart_type("bar")
    await db_session.commit()
    await rerender_report(stored_report.id, session_factory=_fixed_factory(db_session))
    await db_session.refresh(stored_report)
    bar_pdf = stored_report.pdf_data
    bar_pptx = stored_report.pptx_data
    assert bar_pdf is not None
    assert bar_pptx is not None

    stored_report.design = _design_with_chart_type("line")
    await db_session.commit()
    await rerender_report(stored_report.id, session_factory=_fixed_factory(db_session))
    await db_session.refresh(stored_report)
    line_pdf = stored_report.pdf_data
    line_pptx = stored_report.pptx_data

    assert line_pdf != bar_pdf
    # A chart_type override reaches the deck's figure because both artifacts are
    # drawn from the same ``redrawn`` list, which ``_with_chart`` already applied the
    # override to — true before 5b and unchanged by it. 5b changed the deck's
    # *structure*, which now follows the block list at slide granularity.
    assert line_pptx != bar_pptx


def _data_section() -> SectionResult:
    return SectionResult(
        title="Активи Apple",
        source="database",
        sub_question="активи",
        kind="data",
        result=QueryResult(columns=["fy", "value"], rows=[(2024, Decimal("1.0"))]),
    )


def _text_section() -> SectionResult:
    return SectionResult(
        title="Політика відпусток",
        source="documents",
        sub_question="відпустки",
        kind="text",
        prose="Співробітники накопичують відпустку щомісяця.",
    )


async def _stored_report(
    db_session: AsyncSession, *, sections: list[SectionResult]
) -> GeneratedReport:
    """A DONE report carrying an arbitrary section list, for block-list tests.

    Mirrors the ``stored_report`` fixture above, which is pinned to a single
    ``one_data_section`` and so cannot stand in where a test needs a data
    section *and* a text section to tell apart which one a block list kept.
    """
    repo = Repository(name="Vault", llm_provider="gemini", llm_model="gem-1")
    db_session.add(repo)
    await db_session.flush()
    report = GeneratedReport(
        repository_id=repo.id,
        prompt="Revenue for 2024",
        status=ReportStatus.DONE,
    )
    _store_sections(report, title="Звіт", language="uk", sections=sections)
    db_session.add(report)
    await db_session.commit()
    return report


def _deck_titles(blob: bytes) -> list[str]:
    from io import BytesIO

    from pptx import Presentation

    out: list[str] = []
    for slide in Presentation(BytesIO(blob)).slides:
        texts = [s.text_frame.text for s in slide.shapes if s.has_text_frame]
        out.append(texts[0] if texts else "")
    return out


async def test_rerender_restructures_the_deck_not_only_the_document(
    db_session: AsyncSession,
) -> None:
    """A stored block list must reach the deck too — 5a's known divergence, closed.

    Asserted on ``pptx_data`` deliberately: before 5b the deck was drawn from
    ``sections`` regardless of what was stored, so an assertion on ``pdf_data``
    alone passed while the deck ignored the list entirely.
    """
    report = await _stored_report(db_session, sections=[_data_section(), _text_section()])
    report.design = {
        "blocks_version": BLOCKS_VERSION,
        "blocks": [
            {"kind": "cover"},
            {"kind": "heading", "section": 1},
            {"kind": "prose", "section": 1},
        ],
    }
    await db_session.commit()

    await rerender_report(report.id, session_factory=_fixed_factory(db_session))
    await db_session.refresh(report)

    assert report.pptx_data is not None
    titles = _deck_titles(report.pptx_data)
    assert "Активи Apple" not in titles  # the omitted section is gone from the deck
    assert any("Політика" in title for title in titles)


# --- the design canvas's document ---------------------------------------------


async def test_the_preview_is_the_printed_document_in_preview_mode(
    db_session: AsyncSession, stored_report: GeneratedReport
) -> None:
    html = await preview_report_html(stored_report.id, session_factory=_fixed_factory(db_session))

    assert html.startswith("<!doctype html>")
    assert 'data-block-index="0"' in html
    assert 'class="cv-sheet"' in html


async def test_the_preview_draws_the_charts_the_pdf_draws(
    db_session: AsyncSession, stored_report: GeneratedReport
) -> None:
    """A section loaded from storage carries no ``chart_svg`` — the design
    endpoint says so in its own comment. Skip the redraw and the canvas is a
    report with its figures missing, beside a PDF that plainly has them."""
    html = await preview_report_html(stored_report.id, session_factory=_fixed_factory(db_session))

    assert "data:image/svg+xml;base64," in html


async def test_the_preview_touches_neither_the_provider_nor_the_database(
    db_session: AsyncSession,
    monkeypatch: pytest.MonkeyPatch,
    stored_report: GeneratedReport,
) -> None:
    """The same instrument the rerender guarantee uses: stub both so either one
    being called fails the test, rather than asserting that some HTML came back."""

    def explode(*args: Any, **kwargs: Any) -> Any:
        raise AssertionError("the preview must not call out")

    monkeypatch.setattr("contextvault.services.reports.run_validated_query", explode)
    monkeypatch.setattr("contextvault.services.report_llm.generate_text", explode)

    html = await preview_report_html(stored_report.id, session_factory=_fixed_factory(db_session))

    assert "<body" in html


async def test_the_preview_writes_nothing(
    db_session: AsyncSession, stored_report: GeneratedReport
) -> None:
    """Opening the editor must not change what a reader downloads. The editor
    calls this after every debounced edit."""
    stored_report.pdf_data = b"%PDF-the-artifact-of-record"
    stored_report.pptx_data = b"PK-the-deck"
    await db_session.commit()

    await preview_report_html(stored_report.id, session_factory=_fixed_factory(db_session))

    await db_session.refresh(stored_report)
    assert stored_report.pdf_data == b"%PDF-the-artifact-of-record"
    assert stored_report.pptx_data == b"PK-the-deck"


async def test_the_preview_refuses_a_report_with_no_stored_data(
    db_session: AsyncSession, stored_report: GeneratedReport
) -> None:
    stored_report.sections_data = None
    await db_session.commit()

    with pytest.raises(ReportNotRenderable):
        await preview_report_html(stored_report.id, session_factory=_fixed_factory(db_session))


async def test_the_preview_refuses_a_report_that_is_not_done(
    db_session: AsyncSession, stored_report: GeneratedReport
) -> None:
    """Section data is written before the renderers run, so a report that failed
    while rendering can carry a payload. Previewing it would show its owner a
    document for a report they were told had failed."""
    stored_report.status = ReportStatus.FAILED
    await db_session.commit()

    with pytest.raises(ReportNotRenderable):
        await preview_report_html(stored_report.id, session_factory=_fixed_factory(db_session))


async def test_the_preview_and_the_print_read_the_same_stored_design(
    db_session: AsyncSession, stored_report: GeneratedReport
) -> None:
    """One assembly, not two. A block hidden in the design must be absent from
    both, and this is the test that fails if the canvas ever grows its own
    loader."""
    stored_report.design = {
        **(stored_report.design or {}),
        "blocks_version": BLOCKS_VERSION,
        "blocks": [{"kind": "heading", "section": 0, "hidden": True}],
    }
    await db_session.commit()

    html = await preview_report_html(stored_report.id, session_factory=_fixed_factory(db_session))

    assert "<h2" not in html


async def test_an_ambiguous_result_is_retried_once_and_kept_when_the_retry_fails(
    db_session: AsyncSession, monkeypatch: pytest.MonkeyPatch
) -> None:
    """Five rows for 2017 is not a series. The loop must say so and ask again —
    and, like an empty result, keep what it has rather than erroring when the
    second attempt is no better."""
    repo, conn = await _connected_repo(db_session)
    report = GeneratedReport(
        repository_id=repo.id,
        connection_id=conn.id,
        requested_by=None,
        prompt="p",
        plan=_single_db_plan("p", conn),
    )
    db_session.add(report)
    await db_session.flush()

    attempts: list[list[str]] = []

    async def fake_generate(provider: str, api_key: str, model: str, **kwargs: Any) -> ReportQuery:
        attempts.append(list(kwargs["feedback"]))
        return _query("SELECT count(*) AS fiscal_year FROM repositories")

    ambiguous = QueryResult(
        columns=["fiscal_year", "revenue"],
        rows=[(2017, Decimal("1")), (2017, Decimal("2"))],
    )

    async def fake_run(**kwargs: Any) -> QueryResult:
        return ambiguous

    monkeypatch.setattr(reports_service, "generate_report_query", fake_generate)
    monkeypatch.setattr(reports_service, "run_validated_query", fake_run)
    monkeypatch.setattr(
        reports_service.provider_service,  # type: ignore[attr-defined]
        "get_provider_key",
        _fake_key,
    )

    await reports_service.generate_report(db_session, report)

    assert len(attempts) == 2, "the ambiguous result must cost exactly one retry"
    assert attempts[1] and "2017" in attempts[1][0], "the feedback must name the offending value"
    assert report.status is ReportStatus.DONE, "a failed retry delivers, as with an empty result"
    assert report.sections is not None and report.sections[0]["status"] == "done"


async def test_the_measured_period_note_reaches_the_query_prompt(
    db_session: AsyncSession, monkeypatch: pytest.MonkeyPatch
) -> None:
    """The probe and the prompt are wired by three separate hand-offs, and a
    mutation removing the last one survived the whole suite: the note was
    measured and then quietly dropped. Assert it arrives, not just that it exists.
    """
    repo, conn = await _connected_repo(db_session)
    report = GeneratedReport(
        repository_id=repo.id,
        connection_id=conn.id,
        requested_by=None,
        prompt="p",
        plan=_single_db_plan("p", conn),
    )
    db_session.add(report)
    await db_session.flush()

    seen: list[dict[str, str] | None] = []

    async def fake_generate(provider: str, api_key: str, model: str, **kwargs: Any) -> ReportQuery:
        seen.append(kwargs.get("period_notes"))
        return _query("SELECT name FROM repositories")

    # Pinned, not read from the database. Left to the real query this asserts on
    # ambient rows: `repositories` is empty in CI, so the zero-rows retry fires
    # and the count is 2 there and 1 on a developer machine holding ~165 rows.
    async def fake_run(**kwargs: Any) -> QueryResult:
        return QueryResult(columns=["name"], rows=[("Кварта",), ("Президія",)])

    async def fake_probe(**kwargs: Any) -> dict[str, str]:
        return {"repositories": "period lengths present: 90d (2479 rows)"}

    monkeypatch.setattr(reports_service, "generate_report_query", fake_generate)
    monkeypatch.setattr(reports_service, "run_validated_query", fake_run)
    monkeypatch.setattr(reports_service, "probe_period_notes", fake_probe)
    monkeypatch.setattr(
        reports_service.provider_service,  # type: ignore[attr-defined]
        "get_provider_key",
        _fake_key,
    )

    await reports_service.generate_report(db_session, report)

    assert seen == [{"repositories": "period lengths present: 90d (2479 rows)"}]


async def test_an_over_aggregated_query_is_retried_once_and_kept_when_the_retry_fails(
    db_session: AsyncSession, monkeypatch: pytest.MonkeyPatch
) -> None:
    """The SUM is already wrong inside the database, so the rows look perfect.
    The loop must catch it from the SQL, ask once more, and — like an empty or
    ambiguous result — deliver what it has rather than erroring."""
    repo, conn = await _connected_repo(db_session)
    # Widen THIS connection's allow-list only. `_connected_repo` is shared by
    # every test in this file, and its schema gives `repositories` a single
    # `name` column with no period pair — `period_pairs` would return nothing and
    # the check could never fire. Overriding here rather than in the fixture
    # keeps the real probe in the other tests seeing exactly what it sees today.
    conn.exposed_schema = [
        {
            "table": "repositories",
            "description": "",
            "columns": [
                {"name": "name", "description": ""},
                {"name": "period_start", "description": ""},
                {"name": "period_end", "description": ""},
            ],
        }
    ]
    await db_session.flush()
    report = GeneratedReport(
        repository_id=repo.id,
        connection_id=conn.id,
        requested_by=None,
        prompt="p",
        plan=_single_db_plan("p", conn),
    )
    db_session.add(report)
    await db_session.flush()

    attempts: list[list[str]] = []

    async def fake_generate(provider: str, api_key: str, model: str, **kwargs: Any) -> ReportQuery:
        attempts.append(list(kwargs["feedback"]))
        return _query("SELECT name, COUNT(*) AS n FROM repositories GROUP BY 1")

    async def fake_run(**kwargs: Any) -> QueryResult:
        return QueryResult(columns=["name", "n"], rows=[("Кварта", 1), ("Президія", 2)])

    async def fake_probe(**kwargs: Any) -> dict[str, str]:
        return {"repositories": "period lengths present: 90d (2479 rows), 363d (1766 rows)"}

    monkeypatch.setattr(reports_service, "generate_report_query", fake_generate)
    monkeypatch.setattr(reports_service, "run_validated_query", fake_run)
    monkeypatch.setattr(reports_service, "probe_period_notes", fake_probe)
    monkeypatch.setattr(
        reports_service.provider_service,  # type: ignore[attr-defined]
        "get_provider_key",
        _fake_key,
    )

    await reports_service.generate_report(db_session, report)

    assert len(attempts) == 2, "the over-aggregated query must cost exactly one retry"
    assert "repositories" in attempts[1][0], "the feedback must name the table"
    assert "363d" in attempts[1][0], "the feedback must quote the measured durations"
    assert report.status is ReportStatus.DONE
    assert report.sections is not None and report.sections[0]["status"] == "done"
    # The spec made this field a table name rather than a boolean precisely so
    # the admin-visible `sections` audit records *which* table was aggregated.
    assert report.sections[0]["overlap_warning"] == "repositories"
    # The kept section must actually carry the verdict through to the stored
    # render payload — the audit trail above proves only that a section was
    # kept, not that it still knows *why*, and the renderer's caveat and
    # suppressed chart depend on this field surviving the hand-off.
    assert report.sections_data is not None
    stored_sections = report.sections_data["sections"]
    assert stored_sections[0]["overlap_warning"] == "repositories"


async def test_a_frozen_schedule_caveats_its_over_aggregated_sql(
    db_session: AsyncSession, monkeypatch: pytest.MonkeyPatch
) -> None:
    """A frozen schedule re-runs one validated query nightly with no LLM. It
    cannot retry, so it must at least stop claiming the number is sound."""
    repo, conn = await _connected_repo(db_session)
    conn.exposed_schema = [
        {
            "table": "repositories",
            "description": "",
            "columns": [
                {"name": "name", "description": ""},
                {"name": "period_start", "description": ""},
                {"name": "period_end", "description": ""},
            ],
        }
    ]
    await db_session.flush()
    schedule = ReportSchedule(
        repository_id=repo.id,
        connection_id=conn.id,
        owner_id=None,
        prompt="дохід",
        frozen_sql="SELECT name, COUNT(*) AS n FROM repositories GROUP BY 1",
        frozen_chart_spec=ChartSpec(chart_type="bar", x_column="name", y_column="n").model_dump(),
        # NOT NULL on the model; omitting it fails at flush, not at assert.
        run_at_time=time(6, 0),
    )
    db_session.add(schedule)
    await db_session.flush()

    async def fake_run(**kwargs: Any) -> QueryResult:
        return QueryResult(columns=["name", "n"], rows=[("Кварта", 1), ("Президія", 2)])

    async def fake_probe(**kwargs: Any) -> dict[str, str]:
        return {"repositories": "period lengths present: 90d (2479 rows), 363d (1766 rows)"}

    monkeypatch.setattr(reports_service, "run_validated_query", fake_run)
    monkeypatch.setattr(reports_service, "probe_period_notes", fake_probe)

    report = await reports_service.execute_frozen(db_session, schedule=schedule)

    assert report.status is ReportStatus.DONE
    assert report.sections_data is not None
    _, _, sections = load_sections(report.sections_data)
    assert sections[0].overlap_warning == "repositories"


def test_a_rerendered_section_does_not_grow_its_chart_back() -> None:
    """The second chart door. `_with_chart` redraws rather than loads, so a
    re-render would restore the very chart the caveat says was withheld."""
    section = SectionResult(
        title="Дохід",
        source="database",
        sub_question="д",
        kind="data",
        result=QueryResult(columns=["year", "n"], rows=[(2017, 1), (2018, 2)]),
        chart_spec=ChartSpec(chart_type="bar", x_column="year", y_column="n").model_dump(),
        overlap_warning="apple_financials",
    )
    redrawn = reports_service._with_chart(
        section, language="uk", report_title="Звіт", style=resolve_style(None)
    )
    assert redrawn.chart_svg is None


async def _second_connection(
    db_session: AsyncSession, repo: Repository, *, name: str = "Microsoft"
) -> DatabaseConnection:
    """A second database on the same repository (#201), told apart in a prompt.

    Points at the same real database so the guardrailed SQL still executes, but
    carries a distinct allow-list *description*. Guardrails validate table and
    column names only, so the SQL stays valid while the schema block the model is
    shown differs — which is what makes "which database did this section read?"
    observable without stubbing the executor.
    """
    url = make_url(get_settings().database_url)
    conn = DatabaseConnection(
        repository_id=repo.id,
        name=name,
        description=f"{name} filings",
        db_type=DatabaseType.POSTGRES,
        host=url.host or "localhost",
        port=url.port or 5432,
        database=url.database or "",
        username=url.username or "",
        password_encrypted=encrypt(url.password or ""),
        exposed_schema=[
            {
                "table": "repositories",
                "description": f"{name}-SIDE",
                "columns": [{"name": "name", "description": ""}],
            }
        ],
    )
    db_session.add(conn)
    await db_session.flush()
    return conn


class TestSeveralDatabasesInOneReport:
    """#203, the owner's request: Apple in one database, Microsoft in another."""

    @staticmethod
    def _plan(first: DatabaseConnection, second: DatabaseConnection) -> dict[str, Any]:
        return {
            "title": "Порівняння",
            "language": "uk",
            "sections": [
                {
                    "title": "Apple",
                    "source": "database",
                    "databases": [str(first.id)],
                    "sub_question": "apple revenue",
                },
                {
                    "title": "Microsoft",
                    "source": "database",
                    "databases": [str(second.id)],
                    "sub_question": "microsoft revenue",
                },
            ],
        }

    async def test_each_section_runs_on_the_database_its_plan_names(
        self, db_session: AsyncSession, monkeypatch: pytest.MonkeyPatch
    ) -> None:
        """The whole point of the card: two sections of one report, two databases.

        Before #203 the pipeline resolved a single connection for the report, so
        both sections would have been answered from the same one — the Microsoft
        section silently reading Apple's data.
        """
        repo, apple = await _connected_repo(db_session)
        microsoft = await _second_connection(db_session, repo)
        report = GeneratedReport(
            repository_id=repo.id,
            requested_by=None,
            prompt="Порівняй Apple та Microsoft",
            plan=self._plan(apple, microsoft),
        )
        db_session.add(report)
        await db_session.flush()

        seen: list[tuple[str, str]] = []

        async def fake_generate(
            provider: str, api_key: str, model: str, **kwargs: Any
        ) -> ReportQuery:
            seen.append((kwargs["user_prompt"], kwargs["exposed_schema"][0]["description"]))
            return _query("SELECT name FROM repositories")

        monkeypatch.setattr(reports_service, "generate_report_query", fake_generate)
        # Stub the executor: these sections would otherwise run real SQL against
        # whatever earlier runs left committed, so on a clean database every
        # section comes back empty and #185 refuses the whole report. That fails
        # on CI and passes locally — see ``_matching_rows``.
        monkeypatch.setattr(reports_service, "run_validated_query", _matching_rows)

        await reports_service.generate_report(db_session, report)

        assert report.status is ReportStatus.DONE
        # Each sub-question was asked against its OWN database's allow-list.
        assert [prompt for prompt, _ in seen] == ["apple revenue", "microsoft revenue"]
        assert [schema for _, schema in seen] == ["", "Microsoft-SIDE"]

    async def test_the_report_records_the_first_database_it_read(
        self, db_session: AsyncSession, monkeypatch: pytest.MonkeyPatch
    ) -> None:
        """``connection_id`` is written by generation, not guessed at creation.

        Creation cannot know it — the plan is what chooses — so the row is created
        with none and this is where it becomes true. In plan order: "the first
        database this report read" is a fact about the document.
        """
        repo, apple = await _connected_repo(db_session)
        microsoft = await _second_connection(db_session, repo)
        report = GeneratedReport(
            repository_id=repo.id,
            requested_by=None,
            prompt="Порівняй",
            plan=self._plan(apple, microsoft),
        )
        db_session.add(report)
        await db_session.flush()
        assert report.connection_id is None  # creation chose nothing

        async def fake_generate(
            provider: str, api_key: str, model: str, **kwargs: Any
        ) -> ReportQuery:
            return _query("SELECT name FROM repositories")

        monkeypatch.setattr(reports_service, "generate_report_query", fake_generate)
        # Stub the executor: these sections would otherwise run real SQL against
        # whatever earlier runs left committed, so on a clean database every
        # section comes back empty and #185 refuses the whole report. That fails
        # on CI and passes locally — see ``_matching_rows``.
        monkeypatch.setattr(reports_service, "run_validated_query", _matching_rows)

        await reports_service.generate_report(db_session, report)

        assert report.status is ReportStatus.DONE
        assert report.connection_id == apple.id

    async def test_a_section_whose_database_was_deleted_fails_only_that_section(
        self, db_session: AsyncSession, monkeypatch: pytest.MonkeyPatch
    ) -> None:
        """An approved plan is validated, so this means the connection went away
        between approval and generation. Partial success still applies — the other
        section has something to say, and a whole report is not worth losing."""
        repo, apple = await _connected_repo(db_session)
        microsoft = await _second_connection(db_session, repo)
        plan = self._plan(apple, microsoft)
        report = GeneratedReport(
            repository_id=repo.id, requested_by=None, prompt="Порівняй", plan=plan
        )
        db_session.add(report)
        await db_session.flush()
        await db_session.delete(microsoft)
        await db_session.flush()

        async def fake_generate(
            provider: str, api_key: str, model: str, **kwargs: Any
        ) -> ReportQuery:
            return _query("SELECT name FROM repositories")

        monkeypatch.setattr(reports_service, "generate_report_query", fake_generate)
        # Stub the executor: these sections would otherwise run real SQL against
        # whatever earlier runs left committed, so on a clean database every
        # section comes back empty and #185 refuses the whole report. That fails
        # on CI and passes locally — see ``_matching_rows``.
        monkeypatch.setattr(reports_service, "run_validated_query", _matching_rows)

        await reports_service.generate_report(db_session, report)

        assert report.status is ReportStatus.DONE
        assert report.sections is not None
        assert report.sections[0]["status"] == "done"
        assert report.sections[1]["status"] == "error"
        assert "no longer connected" in report.sections[1]["error"]


class TestOneSectionAcrossSeveralDatabases:
    """#204, the card that actually delivers the request: Apple beside Microsoft
    in ONE table, not one section each.

    ``_connected_repo``'s connection is named "Reporting" and the second
    "Microsoft"; those names are what the provenance column prints.
    """

    CONTRACT = ["company", "fiscal_year", "revenue"]

    @staticmethod
    def _plan(first: DatabaseConnection, second: DatabaseConnection) -> dict[str, Any]:
        return {
            "title": "Порівняння",
            "language": "uk",
            "sections": [
                {
                    "title": "Дохід",
                    "source": "database",
                    "databases": [str(first.id), str(second.id)],
                    "columns": TestOneSectionAcrossSeveralDatabases.CONTRACT,
                    "sub_question": "дохід Apple та Microsoft",
                }
            ],
        }

    @staticmethod
    def _stubs(
        monkeypatch: pytest.MonkeyPatch, *, second_rows: list[tuple[Any, ...]] | None = None
    ) -> list[tuple[str, ...]]:
        """Both providers stubbed, keyed on which allow-list the model was shown.

        ``run_validated_query`` opens its own connection, so a test that let it
        run for real would read whatever earlier runs committed — passing locally
        and failing on CI. See ``_matching_rows``.
        """
        contracts: list[tuple[str, ...]] = []

        async def fake_generate(
            provider: str, api_key: str, model: str, **kwargs: Any
        ) -> ReportQuery:
            contracts.append(tuple(kwargs.get("required_columns") or ()))
            side = kwargs["exposed_schema"][0]["description"]
            return ReportQuery(
                sql=f"SELECT name FROM repositories -- {side}",
                chart=ChartSpec(chart_type="bar", x_column="fiscal_year", title="t"),
            )

        async def fake_run(**kwargs: Any) -> QueryResult:
            microsoft = "Microsoft-SIDE" in str(kwargs.get("sql", ""))
            rows = (
                (second_rows if second_rows is not None else [("Microsoft", 2025, 281.7)])
                if microsoft
                else [("Apple", 2025, 416.2)]
            )
            return QueryResult(
                columns=TestOneSectionAcrossSeveralDatabases.CONTRACT, rows=list(rows)
            )

        monkeypatch.setattr(reports_service, "generate_report_query", fake_generate)
        monkeypatch.setattr(reports_service, "run_validated_query", fake_run)
        return contracts

    async def _report(self, db_session: AsyncSession) -> GeneratedReport:
        repo, apple = await _connected_repo(db_session)
        microsoft = await _second_connection(db_session, repo)
        report = GeneratedReport(
            repository_id=repo.id,
            requested_by=None,
            prompt="Порівняй Apple та Microsoft",
            plan=self._plan(apple, microsoft),
        )
        db_session.add(report)
        await db_session.flush()
        return report

    async def test_one_section_holds_the_rows_of_both_databases(
        self, db_session: AsyncSession, monkeypatch: pytest.MonkeyPatch
    ) -> None:
        """Before #204 this needed two sections and the reader did the comparing."""
        report = await self._report(db_session)
        self._stubs(monkeypatch)

        await reports_service.generate_report(db_session, report)

        assert report.status is ReportStatus.DONE
        assert report.sections is not None
        assert len(report.sections) == 1
        assert report.sections_data is not None
        _, _, sections = load_sections(report.sections_data)
        result = sections[0].result
        assert result is not None
        assert [row[1] for row in result.rows] == ["Apple", "Microsoft"]

    async def test_the_merged_table_says_which_database_each_row_came_from(
        self, db_session: AsyncSession, monkeypatch: pytest.MonkeyPatch
    ) -> None:
        report = await self._report(db_session)
        self._stubs(monkeypatch)

        await reports_service.generate_report(db_session, report)

        assert report.sections_data is not None
        _, _, sections = load_sections(report.sections_data)
        result = sections[0].result
        assert result is not None
        assert result.columns[0] == PROVENANCE_COLUMN
        # Plan order, not the alphabetical order ``_load_context`` selects in:
        # "Microsoft" sorts first and still comes second here.
        assert [row[0] for row in result.rows] == ["Reporting", "Microsoft"]

    async def test_every_database_is_told_the_output_contract(
        self, db_session: AsyncSession, monkeypatch: pytest.MonkeyPatch
    ) -> None:
        """Without it the two queries return whatever each database calls its
        columns, and the rows cannot be put in one table at all."""
        report = await self._report(db_session)
        contracts = self._stubs(monkeypatch)

        await reports_service.generate_report(db_session, report)

        assert contracts == [tuple(self.CONTRACT), tuple(self.CONTRACT)]

    async def test_a_database_that_answers_with_nothing_is_named_not_dropped(
        self, db_session: AsyncSession, monkeypatch: pytest.MonkeyPatch
    ) -> None:
        """Partial failure, delivered rather than hidden: the section still
        carries Apple's row, and says Microsoft gave none."""
        report = await self._report(db_session)
        self._stubs(monkeypatch, second_rows=[])

        await reports_service.generate_report(db_session, report)

        assert report.status is ReportStatus.DONE
        assert report.sections_data is not None
        _, _, sections = load_sections(report.sections_data)
        assert sections[0].result is not None
        assert len(sections[0].result.rows) == 1
        assert any("Microsoft" in note for note in sections[0].notes)

    async def test_a_single_database_section_gains_no_provenance_column(
        self, db_session: AsyncSession, monkeypatch: pytest.MonkeyPatch
    ) -> None:
        """The ordinary section is untouched — no merge, no extra column, and no
        contract handed to the query generator."""
        repo, apple = await _connected_repo(db_session)
        report = GeneratedReport(
            repository_id=repo.id,
            requested_by=None,
            prompt="Дохід",
            plan=_single_db_plan("apple revenue", apple),
        )
        db_session.add(report)
        await db_session.flush()
        contracts = self._stubs(monkeypatch)

        await reports_service.generate_report(db_session, report)

        assert report.status is ReportStatus.DONE
        assert report.sections_data is not None
        _, _, sections = load_sections(report.sections_data)
        assert sections[0].result is not None
        assert sections[0].result.columns[0] != PROVENANCE_COLUMN
        assert contracts == [()]

    async def test_a_database_deleted_after_approval_is_described_not_printed_as_an_id(
        self, db_session: AsyncSession, monkeypatch: pytest.MonkeyPatch
    ) -> None:
        """A plan approved before a connection was deleted still generates. Only
        the id survives that deletion, and a UUID in the reader's note tells them
        less than nothing — the same reason the plan editor's chip prints none."""
        repo, apple = await _connected_repo(db_session)
        gone = uuid.uuid4()
        report = GeneratedReport(
            repository_id=repo.id,
            requested_by=None,
            prompt="Порівняй",
            plan={
                "title": "Порівняння",
                "language": "uk",
                "sections": [
                    {
                        "title": "Дохід",
                        "source": "database",
                        "databases": [str(apple.id), str(gone)],
                        "columns": self.CONTRACT,
                        "sub_question": "дохід",
                    }
                ],
            },
        )
        db_session.add(report)
        await db_session.flush()
        self._stubs(monkeypatch)

        await reports_service.generate_report(db_session, report)

        assert report.status is ReportStatus.DONE
        assert report.sections_data is not None
        _, _, sections = load_sections(report.sections_data)
        note = " ".join(sections[0].notes)
        assert str(gone) not in note
        assert "від'єднано" in note

    async def test_the_section_records_every_database_it_read(
        self, db_session: AsyncSession, monkeypatch: pytest.MonkeyPatch
    ) -> None:
        """#205: the renderer captions a merged figure from this, rather than
        reverse-engineering provenance out of the result's magic column — which
        would give a chart-only design nothing at all."""
        report = await self._report(db_session)
        self._stubs(monkeypatch)

        await reports_service.generate_report(db_session, report)

        assert report.sections_data is not None
        _, _, sections = load_sections(report.sections_data)
        assert sections[0].databases == ["Reporting", "Microsoft"]

    async def test_a_database_that_gave_nothing_is_still_recorded_as_read(
        self, db_session: AsyncSession, monkeypatch: pytest.MonkeyPatch
    ) -> None:
        """ "Where could these figures have come from" is what a caption answers,
        and a database that was asked and was silent is part of that answer."""
        report = await self._report(db_session)
        self._stubs(monkeypatch, second_rows=[])

        await reports_service.generate_report(db_session, report)

        assert report.sections_data is not None
        _, _, sections = load_sections(report.sections_data)
        assert sections[0].databases == ["Reporting", "Microsoft"]

    async def test_a_single_database_section_records_its_one_database(
        self, db_session: AsyncSession, monkeypatch: pytest.MonkeyPatch
    ) -> None:
        """Set wherever it is known. Nothing renders it — the caption needs more
        than one — but a field populated only sometimes is the confusing one."""
        repo, apple = await _connected_repo(db_session)
        report = GeneratedReport(
            repository_id=repo.id,
            requested_by=None,
            prompt="Дохід",
            plan=_single_db_plan("apple revenue", apple),
        )
        db_session.add(report)
        await db_session.flush()
        self._stubs(monkeypatch)

        await reports_service.generate_report(db_session, report)

        assert report.sections_data is not None
        _, _, sections = load_sections(report.sections_data)
        assert sections[0].databases == ["Reporting"]
